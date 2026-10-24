"""擂台赛端到端测试（使用临时数据库，不影响正式数据）

运行：cd backend && python test_arena_e2e.py
"""
import os
import sys
import json
import pathlib
import tempfile

_tmp_db = pathlib.Path(tempfile.gettempdir()) / "arena_e2e_test.db"
if _tmp_db.exists():
    _tmp_db.unlink()
os.environ["DATABASE_URL"] = "sqlite:///" + _tmp_db.as_posix()

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient

from app.main import app
from app.database import SessionLocal
from app import models

PASSED = []
FAILED = []


def check(name, cond, extra=""):
    if cond:
        PASSED.append(name)
        print("  [PASS] " + name)
    else:
        FAILED.append((name + " " + str(extra)).strip())
        print("  [FAIL] " + name + " " + str(extra))


def drain(ws):
    """发送 ping 并读取直到 pong，返回期间收到的所有消息"""
    ws.send_json({"type": "ping"})
    msgs = []
    while True:
        m = ws.receive_json()
        msgs.append(m)
        if m["type"] == "pong":
            return msgs


def seed_questions():
    """插入六道题（客观题、多选题各若干 + 一道主观题），供擂台赛引用"""
    db = SessionLocal()
    try:
        q1 = models.Question(
            type="single",
            content="<p>心率正常范围是多少？</p>",
            options=json.dumps({"A": "60-100", "B": "40-60", "C": "100-120", "D": "30-40"}),
            answer="A",
            analysis="正常心率为60-100次/分",
            category="生理学",
        )
        q2 = models.Question(
            type="multiple",
            content="<p>哪些属于基本生命体征？</p>",
            options=json.dumps({"A": "体温", "B": "脉搏", "C": "呼吸", "D": "血压", "E": "身高"}),
            answer="A,B,C,D",
            analysis="基本生命体征包括体温、脉搏、呼吸、血压",
            category="生理学",
        )
        q3 = models.Question(
            type="multiple",
            content="<p>请简述心力衰竭的代偿机制。</p>",
            options="{}",
            answer="",
            analysis="代偿机制包括交感兴奋、心肌肥厚等",
            category="内科",
        )
        # 三道标准答案相同的多选题，用于验证集合精确匹配（少选/多选均判错）
        q4 = models.Question(
            type="multiple",
            content="<p>多选题1：哪些属于基本生命体征？</p>",
            options=json.dumps({"A": "体温", "B": "脉搏", "C": "呼吸", "D": "血压", "E": "身高"}),
            answer="A,B,C,D",
            analysis="基本生命体征包括体温、脉搏、呼吸、血压",
            category="生理学",
        )
        q5 = models.Question(
            type="multiple",
            content="<p>多选题2：哪些属于基本生命体征？</p>",
            options=json.dumps({"A": "体温", "B": "脉搏", "C": "呼吸", "D": "血压", "E": "身高"}),
            answer="A,B,C,D",
            analysis="基本生命体征包括体温、脉搏、呼吸、血压",
            category="生理学",
        )
        q6 = models.Question(
            type="multiple",
            content="<p>多选题3：哪些属于基本生命体征？</p>",
            options=json.dumps({"A": "体温", "B": "脉搏", "C": "呼吸", "D": "血压", "E": "身高"}),
            answer="A,B,C,D",
            analysis="基本生命体征包括体温、脉搏、呼吸、血压",
            category="生理学",
        )
        db.add_all([q1, q2, q3, q4, q5, q6])
        db.commit()
        return q1.id, q2.id, q3.id, q4.id, q5.id, q6.id
    finally:
        db.close()


def main():
    # 使用上下文管理器以触发 startup 事件（创建管理员账号）
    client = TestClient(app)
    client.__enter__()

    print("")
    print("== 1. 管理员登录 ==")
    r = client.post("/api/auth/login", json={"employee_id": "admin", "password": "admin123"})
    check("管理员登录成功", r.status_code == 200, r.text[:200])
    token = r.json()["access_token"]
    H = {"Authorization": "Bearer " + token}

    q1_id, q2_id, q3_id, q4_id, q5_id, q6_id = seed_questions()

    print("")
    print("== 2. 创建擂台赛 ==")
    r = client.post("/api/arenas/", headers=H, json={
        "title": "测试擂台赛",
        "settings": {"required_score": 10, "rush_wrong_penalty": 5,
                     "required_wrong_penalty": False},
        "team_names": ["红队", "蓝队"],
        "team_colors": ["#F56C6C", "#409EFF"],
    })
    check("创建擂台赛", r.status_code == 200, r.text[:300])
    arena = r.json()
    aid = arena["id"]
    check("自动生成6位邀请码", len(arena.get("invite_code") or "") == 6, str(arena.get("invite_code")))
    check("返回邀请码有效期", arena.get("invite_expire") is not None)
    check("队伍数量为2", arena.get("team_count") == 2, str(arena.get("team_count")))

    r = client.get("/api/arenas/" + str(aid) + "/teams")
    teams = r.json()
    check("获取队伍列表", r.status_code == 200 and len(teams) == 2)
    check("队伍颜色生效", teams[0]["color"] == "#F56C6C", teams[0]["color"])
    t1 = teams[0]["id"]
    t2 = teams[1]["id"]

    print("")
    print("== 3. 添加擂台赛题目 ==")
    r = client.post("/api/arenas/" + str(aid) + "/questions", headers=H, json=[
        {"question_id": q1_id, "q_type": "required", "order_num": 0, "score": 10},
        {"question_id": q2_id, "q_type": "rush", "order_num": 1, "score": 20},
    ])
    check("添加题目", r.status_code == 200 and len(r.json()) == 2, r.text[:300])

    r = client.get("/api/arenas/" + str(aid) + "/questions/public")
    pub = r.json()
    check("公开题目列表不含答案", r.status_code == 200 and all("answer" not in q for q in pub))
    check("公开题目不含解析", all("analysis" not in q for q in pub))
    aq1 = pub[0]["id"]
    aq2 = pub[1]["id"]

    print("")
    print("== 4. 邀请码加入 ==")
    r = client.get("/api/arenas/invite/" + arena["invite_code"])
    check("通过邀请码查询擂台赛", r.status_code == 200 and r.json()["id"] == aid, r.text[:200])

    r = client.post("/api/arenas/players/join", json={
        "team_id": t1, "device_id": "dev-e2e-1", "player_name": "选手甲"})
    check("选手加入并获得令牌", r.status_code == 200 and bool(r.json().get("token")), r.text[:200])
    p1_token = r.json()["token"]

    r = client.post("/api/arenas/players/join", json={
        "team_id": t2, "device_id": "dev-e2e-2", "player_name": "选手乙"})
    check("第二名选手加入", r.status_code == 200 and bool(r.json().get("token")), r.text[:200])
    p2_token = r.json()["token"]

    r = client.post("/api/arenas/players/join", json={
        "team_id": t1, "device_id": "dev-e2e-1", "player_name": "选手甲"})
    check("同一设备重复加入返回原记录", r.status_code == 200 and r.json().get("id") is not None)

    print("")
    print("== 5. WebSocket：主持人 + 选手 ==")
    with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=host&token=" + token) as host:
        check("主持人连接成功", host.receive_json()["type"] == "connected")

        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p1_token) as p1:
            check("选手连接成功", p1.receive_json()["type"] == "connected")

            host.send_json({"type": "host_control",
                            "payload": {"action": "show_question", "question_id": aq1}})
            msg = p1.receive_json()
            check("选手收到题目", msg["type"] == "question_shown", json.dumps(msg, ensure_ascii=False)[:200])
            check("题目不含答案", "answer" not in msg["payload"]["question"])

            p1.send_json({"type": "answer_submit",
                          "payload": {"question_id": aq1, "answer": "A", "time_used": 8}})
            ack = p1.receive_json()
            check("选手收到提交回执", ack["type"] == "answer_received", json.dumps(ack, ensure_ascii=False))
            answer_id = ack["payload"]["answer_id"]

            hmsg = host.receive_json()
            while hmsg["type"] != "answer_submitted":
                hmsg = host.receive_json()
            check("主持人收到待判定提交", hmsg["payload"]["answer_id"] == answer_id)
            check("提交时分数为0（不自动判分）", hmsg["payload"]["score"] == 0)

            host.send_json({"type": "host_control",
                            "payload": {"action": "judge_answer", "answer_id": answer_id,
                                        "is_correct": True}})
            j = p1.receive_json()
            check("选手收到判定结果", j["type"] == "judge_result", json.dumps(j, ensure_ascii=False))
            check("判定为正确", j["payload"]["is_correct"] is True)
            check("得分等于题目分值10", j["payload"]["score"] == 10, str(j["payload"].get("score")))
            check("队伍总分同步为10", j["payload"]["team_score"] == 10, str(j["payload"].get("team_score")))

            host.send_json({"type": "host_control",
                            "payload": {"action": "judge_answer", "answer_id": answer_id,
                                        "is_correct": True}})
            j2 = p1.receive_json()
            while j2["type"] != "judge_result":
                j2 = p1.receive_json()
            check("重复判定不重复加分", j2["payload"]["team_score"] == 10, str(j2["payload"].get("team_score")))

            host.send_json({"type": "host_control",
                            "payload": {"action": "judge_answer", "answer_id": answer_id,
                                        "is_correct": False}})
            j3 = p1.receive_json()
            while j3["type"] != "judge_result":
                j3 = p1.receive_json()
            check("改判后分数回退", j3["payload"]["team_score"] == 0, str(j3["payload"].get("team_score")))

    print("")
    print("== 6. 求助（服务端排除错误选项，仅发给本队）==")
    with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=host&token=" + token) as host:
        host.receive_json()
        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p1_token) as p1:
            p1.receive_json()
            with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p2_token) as p2:
                p2.receive_json()
                p1.send_json({"type": "help_request",
                              "payload": {"help_type": "eliminate", "question_id": aq1}})
                h = p1.receive_json()
                check("求助成功", h["type"] == "help_result" and h["payload"]["success"] is True,
                      json.dumps(h, ensure_ascii=False)[:200])
                excluded = h["payload"]["excluded_options"]
                check("排除2个错误选项", len(excluded) == 2, str(excluded))
                check("排除项不含正确答案A", "A" not in [e.upper() for e in excluded], str(excluded))

                other_msgs = drain(p2)
                check("求助结果未泄露给其他队伍",
                      not any(m["type"] == "help_result" for m in other_msgs),
                      str([m["type"] for m in other_msgs]))

    print("")
    print("== 7. 抢答仲裁 ==")
    with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=host&token=" + token) as host:
        host.receive_json()
        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p1_token) as p1:
            p1.receive_json()
            with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p2_token) as p2:
                p2.receive_json()
                host.send_json({"type": "host_control",
                                "payload": {"action": "start_rush", "question_id": aq2}})
                m = p1.receive_json()
                check("选手收到抢答开始", m["type"] == "rush_started", json.dumps(m, ensure_ascii=False))

                p1.send_json({"type": "rush_answer", "payload": {"question_id": aq2}})
                p2.send_json({"type": "rush_answer", "payload": {"question_id": aq2}})
                all_msgs = drain(p1) + drain(p2)
                wins = [m for m in all_msgs if m["type"] == "rush_success"]
                lost = [m for m in all_msgs if m["type"] == "rush_lost"]
                winners = {m["payload"]["team_id"] for m in wins}
                check("只有一支队伍抢到答题权", len(winners) == 1, str(winners))
                check("另一支队伍收到抢答失败", len(lost) == 1, str(len(lost)))

    print("")
    print("== 8. 淘汰 / 复活 ==")
    r = client.post("/api/arenas/" + str(aid) + "/eliminate", headers=H, json={"team_ids": [t2]})
    check("淘汰队伍", r.status_code == 200, r.text[:200])
    r = client.get("/api/arenas/" + str(aid) + "/teams")
    st = {t["id"]: t["status"] for t in r.json()}
    check("被淘汰队伍状态更新", st[t2] == "eliminated", str(st))

    r = client.post("/api/arenas/" + str(aid) + "/resurrect", headers=H, json={"team_id": t2})
    check("复活队伍", r.status_code == 200, r.text[:200])
    r = client.get("/api/arenas/" + str(aid) + "/teams")
    st = {t["id"]: t["status"] for t in r.json()}
    check("复活后状态为resurrected", st[t2] == "resurrected", str(st))

    print("")
    print("== 9. 排名 ==")
    r = client.get("/api/arenas/" + str(aid) + "/standings")
    standings = r.json()
    check("排名接口可用", r.status_code == 200 and len(standings["teams"]) == 2, r.text[:200])
    check("排名含名次", all("rank" in t for t in standings["teams"]))

    print("")
    print("== 10. 导出 Excel ==")
    r = client.get("/api/arenas/" + str(aid) + "/export", headers=H)
    check("导出返回 xlsx", r.status_code == 200 and r.content[:2] == b"PK", str(r.status_code))
    out = pathlib.Path(tempfile.gettempdir()) / "arena_e2e_export.xlsx"
    out.write_bytes(r.content)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(out)
        check("Excel 含3个工作表", len(wb.sheetnames) == 3, str(wb.sheetnames))
    except Exception as e:
        check("Excel 可正常打开", False, str(e))

    print("")
    print("== 11. 邀请码刷新 ==")
    r = client.post("/api/arenas/" + str(aid) + "/refresh-invite", headers=H)
    check("刷新邀请码", r.status_code == 200 and r.json()["invite_code"] != arena["invite_code"],
          r.text[:200])

    print("")
    print("== 12. 权限校验 ==")
    r = client.post("/api/arenas/", json={"title": "无权限"})
    check("未登录无法创建擂台赛", r.status_code in (401, 403), str(r.status_code))
    r = client.get("/api/arenas/" + str(aid) + "/questions")
    check("完整题目列表仅管理员可见", r.status_code in (401, 403), str(r.status_code))

    print("")
    print("== 13. 超时自动判错 ==")
    with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=host&token=" + token) as host:
        host.receive_json()
        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p2_token) as p1:
            p1.receive_json()
            p1.send_json({"type": "answer_submit",
                          "payload": {"question_id": aq1, "answer": "", "time_used": 15,
                                      "timeout": True}})
            msgs = drain(p1)
            types = [m["type"] for m in msgs]
            check("超时提交收到判定结果", "judge_result" in types, str(types))
            jr = [m for m in msgs if m["type"] == "judge_result"]
            if jr:
                check("超时判定为错误", jr[0]["payload"]["is_correct"] is False)
                check("超时不得分（必答题）", jr[0]["payload"]["score"] == 0,
                      str(jr[0]["payload"].get("score")))

    print("")
    print("== 14. 综合分析题由主持人给分 ==")
    r = client.post("/api/arenas/" + str(aid) + "/questions", headers=H, json=[
        {"question_id": q3_id, "q_type": "analysis", "order_num": 2, "score": 100},
    ])
    analysis_list = [q for q in r.json() if q["q_type"] == "analysis"]
    check("添加分析题", r.status_code == 200 and len(analysis_list) == 1, r.text[:200])
    aq3 = analysis_list[0]["id"]
    with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=host&token=" + token) as host:
        host.receive_json()
        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p2_token) as p2:
            p2.receive_json()
            p2.send_json({"type": "answer_submit",
                          "payload": {"question_id": aq3, "answer": "答案要点", "time_used": 30}})
            msgs = drain(p2)
            ack = [m for m in msgs if m["type"] == "answer_received"]
            check("分析题提交成功", len(ack) == 1, str([m["type"] for m in msgs]))
            answer_id = ack[0]["payload"]["answer_id"]
            host.send_json({"type": "host_control",
                            "payload": {"action": "judge_answer", "answer_id": answer_id,
                                        "is_correct": True, "score": 80}})
            msgs2 = drain(p2)
            jr = [m for m in msgs2 if m["type"] == "judge_result"]
            check("分析题收到判定", len(jr) == 1, str([m["type"] for m in msgs2]))
            if jr:
                check("分析题按主持人给分", jr[0]["payload"]["score"] == 80,
                      str(jr[0]["payload"].get("score")))

    print("")
    print("== 15. 主持人断线自动暂停 ==")
    client.put("/api/arenas/" + str(aid), headers=H, json={"status": "active"})
    with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p1_token) as p1:
        p1.receive_json()
        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=host&token=" + token) as host:
            host.receive_json()
            host.close()
        msgs = drain(p1)
        check("选手收到主持人断线通知",
              any(m["type"] == "host_disconnected" for m in msgs),
              str([m["type"] for m in msgs]))
    r = client.get("/api/arenas/" + str(aid))
    check("比赛自动转为暂停", r.json()["status"] == "paused", r.json()["status"])

    print("")
    print("== 16. 断线重连恢复现场 ==")
    client.put("/api/arenas/" + str(aid), headers=H, json={"status": "active"})
    r = client.get("/api/arenas/" + str(aid) + "/questions", headers=H)
    aq_list = r.json()
    # 用尚未作答过的题目（抢答题），保证存在“已提交未判定”的记录
    target_q = aq_list[1]["id"]
    # 主持人展示题目
    with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=host&token=" + token) as host:
        host.receive_json()
        host.send_json({"type": "host_control",
                        "payload": {"action": "show_question", "question_id": target_q}})
        drain(host)
        # 选手提交答案后断线
        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p1_token) as p1:
            p1.receive_json()
            p1.send_json({"type": "answer_submit",
                          "payload": {"question_id": target_q, "answer": "A", "time_used": 8}})
            drain(p1)
        # 选手重连 -> sync_state 应带回当前题目与本队作答状态
        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=player&token=" + p1_token) as p1b:
            p1b.receive_json()
            p1b.send_json({"type": "sync_state", "payload": {}})
            msgs = drain(p1b)
            ss = [m for m in msgs if m["type"] == "state_sync"]
            check("重连收到 state_sync", len(ss) == 1, str([m["type"] for m in msgs]))
            if ss:
                pl = ss[0]["payload"]
                check("重连恢复当前题目", (pl.get("current_question") or {}).get("id") == target_q,
                      str(pl.get("current_question")))
                check("重连恢复作答状态", (pl.get("my_answer_state") or {}).get("question_id") == target_q,
                      str(pl.get("my_answer_state")))
                check("重连恢复队伍分数", isinstance(pl.get("teams"), list) and len(pl["teams"]) > 0)

        # 主持人重连 -> 应带回待判定列表
        with client.websocket_connect("/ws/arenas/" + str(aid) + "?role=host&token=" + token) as host2:
            host2.receive_json()
            host2.send_json({"type": "sync_state", "payload": {}})
            msgs = drain(host2)
            ss = [m for m in msgs if m["type"] == "state_sync"]
            check("主持人重连收到 state_sync", len(ss) == 1, str([m["type"] for m in msgs]))
            if ss:
                pl = ss[0]["payload"]
                pending = pl.get("pending_answers") or []
                check("主持人重连恢复待判定列表", len(pending) == 1, str(pending))
                check("主持人重连恢复当前题目", (pl.get("current_question") or {}).get("id") == target_q)

    print("")
    print("== 17. 时间戳带时区（UTC）==")
    r = client.get("/api/arenas/" + str(aid))
    arena_json = r.json()
    check("created_at 带时区标记",
          str(arena_json.get("created_at", "")).endswith("Z") or "+00:00" in str(arena_json.get("created_at")),
          str(arena_json.get("created_at")))
    check("invite_expire 带时区标记",
          str(arena_json.get("invite_expire", "")).endswith("Z") or "+00:00" in str(arena_json.get("invite_expire")),
          str(arena_json.get("invite_expire")))
    r = client.get("/api/arenas/" + str(aid) + "/logs")
    logs = r.json()
    if logs:
        check("日志时间戳带时区标记",
              str(logs[0].get("timestamp", "")).endswith("Z") or "+00:00" in str(logs[0].get("timestamp")),
              str(logs[0].get("timestamp")))

    print("")
    print("== 18. 多选题自动判定（集合精确匹配）==")
    # 新增三道多选题（标准答案 A,B,C,D），用于覆盖集合精确匹配的三种情况
    r = client.post("/api/arenas/" + str(aid) + "/questions", headers=H, json=[
        {"question_id": q4_id, "q_type": "rush", "order_num": 3, "score": 20},
        {"question_id": q5_id, "q_type": "rush", "order_num": 4, "score": 20},
        {"question_id": q6_id, "q_type": "rush", "order_num": 5, "score": 20},
    ])
    added = [q for q in r.json() if q["q_type"] == "rush"][-3:]
    check("新增3道多选题", r.status_code == 200 and len(added) == 3, r.text[:300])
    aq_m1, aq_m2, aq_m3 = added[0]["id"], added[1]["id"], added[2]["id"]

    r = client.post("/api/arenas/answers", json={
        "team_id": t1, "question_id": aq_m1, "answer": "A,B,C,D", "time_used": 6})
    check("提交多选答案", r.status_code == 200, r.text[:200])
    full_answer_id = r.json()["id"]
    r = client.post("/api/arenas/answers/" + str(full_answer_id) + "/judge", headers=H,
                    json={"is_correct": None})
    check("完全一致的多选答案判定为正确", r.json().get("is_correct") is True, r.text[:200])
    check("多选题得分等于分值20", r.json().get("score") == 20, str(r.json().get("score")))

    r = client.post("/api/arenas/answers", json={
        "team_id": t2, "question_id": aq_m2, "answer": "A,B,C", "time_used": 6})
    partial_answer_id = r.json()["id"]
    r = client.post("/api/arenas/answers/" + str(partial_answer_id) + "/judge", headers=H,
                    json={"is_correct": None})
    check("缺少选项的多选答案判定为错误", r.json().get("is_correct") is False, r.text[:200])

    r = client.post("/api/arenas/answers", json={
        "team_id": t2, "question_id": aq_m3, "answer": "A,B,C,D,E", "time_used": 6})
    extra_answer_id = r.json()["id"]
    r = client.post("/api/arenas/answers/" + str(extra_answer_id) + "/judge", headers=H,
                    json={"is_correct": None})
    check("多选多余选项判定为错误", r.json().get("is_correct") is False, r.text[:200])

    print("")
    print("== 19. 主持人判定幂等（重复/改判不漂移分数）==")
    r = client.get("/api/arenas/" + str(aid) + "/teams")
    score_before = [t for t in r.json() if t["id"] == t1][0]["score"]
    r = client.post("/api/arenas/answers/" + str(full_answer_id) + "/judge", headers=H,
                    json={"is_correct": True})
    check("重复判定为正确不重复加分", r.json().get("score") == 20, str(r.json().get("score")))
    r = client.post("/api/arenas/answers/" + str(full_answer_id) + "/judge", headers=H,
                    json={"is_correct": False})
    check("改判为错误后分值回退", r.json().get("score") == -5, str(r.json().get("score")))
    r = client.get("/api/arenas/" + str(aid) + "/teams")
    t1_score = [t for t in r.json() if t["id"] == t1][0]["score"]
    check("队伍总分按差额同步", t1_score == score_before - 25, str((score_before, t1_score)))

    print("")
    print("=" * 50)
    print("通过 " + str(len(PASSED)) + " 项，失败 " + str(len(FAILED)) + " 项")
    if FAILED:
        print("失败项：")
        for f in FAILED:
            print("  - " + f)
        sys.exit(1)
    print("全部通过")


if __name__ == "__main__":
    main()
