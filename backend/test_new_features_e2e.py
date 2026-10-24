"""新增功能端到端测试（临时数据库，不影响正式数据）

覆盖：
  1. 普通考试成绩 Excel 导出（总成绩表 / 答题明细 / 统计表）
  2. 擂台赛自动判分（单条 auto_judge_answer / 批量 auto_judge_all）
  3. 题目分类筛选（分类 / 题型）

运行：cd backend && python test_new_features_e2e.py
"""
import os
import sys
import json
import pathlib
import tempfile
import io

_tmp_db = pathlib.Path(tempfile.gettempdir()) / "new_features_e2e_test.db"
if _tmp_db.exists():
    _tmp_db.unlink()
os.environ["DATABASE_URL"] = "sqlite:///" + _tmp_db.as_posix()

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient

from app.main import app
from app.database import SessionLocal
from app import models

PASSED = []
FAILED = []


def check(name, cond, extra=""):
    if cond:
        PASSED.append(name)
        print("  [PASS] " + name)
    else:
        FAILED.append((name + " " + str(extra)).strip())
        print("  [FAIL] " + name + " " + str(extra))


def drain(ws):
    ws.send_json({"type": "ping"})
    msgs = []
    while True:
        m = ws.receive_json()
        msgs.append(m)
        if m["type"] == "pong":
            return msgs


def recv_until(ws, want, max_msgs=60):
    """持续接收直到出现指定 type 的消息（WS 并发下 pong 可能先于业务消息到达）"""
    got = []
    for _ in range(max_msgs):
        m = ws.receive_json()
        got.append(m)
        if m.get("type") in want:
            return m, got
    return None, got


def seed_questions():
    """插入形态学单选/判断/多选题各若干"""
    db = SessionLocal()
    try:
        q1 = models.Question(
            type="single",
            content="<p>红细胞呈双凹圆盘状的是？</p>",
            options=json.dumps({"A": "红细胞", "B": "白细胞", "C": "血小板", "D": "淋巴细胞"}),
            answer="A",
            analysis="红细胞为双凹圆盘状",
            category="形态学/血液形态学",
        )
        q2 = models.Question(
            type="single",
            content="<p>肾小球滤过膜最外层是？</p>",
            options=json.dumps({"A": "内皮细胞", "B": "基膜", "C": "上皮细胞"}),
            answer="C",
            analysis="上皮细胞在滤过膜最外层",
            category="形态学/泌尿形态学",
        )
        q3 = models.Question(
            type="judgment",
            content="<p>判断：红细胞具有中央淡染区。</p>",
            options=json.dumps({"A": "", "B": "", "C": "", "D": ""}),
            answer="正确",
            analysis="",
            category="形态学/血液形态学",
        )
        q4 = models.Question(
            type="multiple",
            content="<p>哪些属于有核细胞？</p>",
            options=json.dumps({"A": "白细胞", "B": "红细胞", "C": "血小板"}),
            answer="A",
            analysis="只有白细胞有核",
            category="形态学/血液形态学",
        )
        # 专用判断题：选项无文本，无法自动判定（需人工评分）
        q5 = models.Question(
            type="judgment",
            content="<p>判断：血小板具有中央颗粒区。</p>",
            options=json.dumps({"A": "", "B": "", "C": "", "D": ""}),
            answer="正确",
            analysis="",
            category="形态学/血液形态学",
        )
        db.add_all([q1, q2, q3, q4, q5])
        db.commit()
        return q1.id, q2.id, q3.id, q4.id, q5.id
    finally:
        db.close()


def seed_user(employee_id, name, dept):
    db = SessionLocal()
    try:
        u = models.User(
            employee_id=employee_id,
            name=name,
            password_hash="x",
            dept=dept,
            role="医生",
            is_admin=False,
        )
        db.add(u)
        db.commit()
        return u.id
    finally:
        db.close()


def main():
    client = TestClient(app)
    client.__enter__()

    print("")
    print("== 1. 登录与准备数据 ==")
    r = client.post("/api/auth/login", json={"employee_id": "admin", "password": "admin123"})
    check("管理员登录成功", r.status_code == 200, r.text[:200])
    token = r.json()["access_token"]
    H = {"Authorization": "Bearer " + token}

    q1, q2, q3, q4, q5 = seed_questions()
    u1 = seed_user("E001", "张三", "检验科")
    u2 = seed_user("E002", "李四", "病理科")

    print("")
    print("== 2. 题目分类筛选 ==")
    r = client.get("/api/questions", headers=H, params={"category": "形态学/血液形态学"})
    got_ids = {q["id"] for q in r.json()["items"]} if r.status_code == 200 else set()
    check("按子分类筛选", r.status_code == 200 and {q1, q3, q4} <= got_ids, str(got_ids))
    r = client.get("/api/questions", headers=H, params={"category": "形态学"})
    got_ids = {q["id"] for q in r.json()["items"]} if r.status_code == 200 else set()
    check("按父分类筛选包含子分类", r.status_code == 200 and {q1, q2, q3, q4} <= got_ids, str(got_ids))
    r = client.get("/api/questions", headers=H, params={"type": "judgment"})
    got_ids = {q["id"] for q in r.json()["items"]} if r.status_code == 200 else set()
    check("按题型筛选", r.status_code == 200 and q3 in got_ids and q1 not in got_ids, str(got_ids))
    r = client.get("/api/questions", headers=H, params={"category": "形态学", "type": "multiple"})
    got_ids = {q["id"] for q in r.json()["items"]} if r.status_code == 200 else set()
    check("分类 + 题型组合筛选", r.status_code == 200 and got_ids == {q4}, str(got_ids))
    r = client.get("/api/questions", headers=H, params={"limit": 1})
    check("分页返回 total 正确", r.status_code == 200 and r.json()["total"] >= 4, r.text[:200])

    print("")
    print("== 3. 普通考试成绩 Excel 导出 ==")
    r = client.post("/api/exams", headers=H, json={
        "title": "形态学考核",
        "duration": 60,
        "pass_score": 60,
        "single_score": 10,
        "multiple_score": 20,
        "judgment_score": 5,
        "question_ids": [q1, q2, q3, q4],
    })
    check("创建考试", r.status_code == 200, r.text[:300])
    exam_id = r.json()["id"]
    db = SessionLocal()
    try:
        exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
        check("考试总分计算正确", exam is not None and exam.total_score == 45,
              str(exam.total_score if exam else None))
    finally:
        db.close()

    db = SessionLocal()
    try:
        db.add(models.ExamRecord(
            exam_id=exam_id, user_id=u1,
            answers=json.dumps({str(q1): "A", str(q2): "C", str(q3): "正确", str(q4): "A"}),
            score=45.0, total_score=45.0,
        ))
        db.add(models.ExamRecord(
            exam_id=exam_id, user_id=u2,
            answers=json.dumps({str(q1): "B", str(q2): "A", str(q3): "错误", str(q4): "A"}),
            score=20.0, total_score=45.0,
        ))
        db.commit()
    finally:
        db.close()

    r = client.get(f"/api/exams/{exam_id}/export", headers=H)
    check("导出接口返回 200", r.status_code == 200, r.text[:300])
    ctype = r.headers.get("content-type", "")
    cdisp = r.headers.get("content-disposition", "")
    check("Content-Type 为 xlsx", "spreadsheetml" in ctype, ctype)
    check("Content-Disposition 含 UTF-8 文件名", "filename*=UTF-8''" in cdisp, cdisp)
    check("文件非空", len(r.content) > 3000, len(r.content))

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    check("包含 3 个工作表", wb.sheetnames == ["总成绩表", "答题明细", "统计表"], str(wb.sheetnames))

    ws1 = wb["总成绩表"]
    head1 = [c.value for c in ws1[1]]
    expect1 = ["排名", "工号", "姓名", "部门", "得分", "总分", "百分制", "是否及格", "提交时间"]
    check("总成绩表表头正确", head1 == expect1, str(head1))
    check("总成绩表有 2 条记录", ws1.max_row == 3, ws1.max_row)
    rows = [[c.value for c in row] for row in ws1.iter_rows(min_row=2)]
    check("满分考生排第一", rows[0][0] == 1 and rows[0][6] == 100.0, str(rows[0]))
    check("低分考生排第二", rows[1][0] == 2 and rows[1][6] == 44.4, str(rows[1]))
    check("及格判定正确", rows[0][7] == "及格" and rows[1][7] == "不及格", str(rows[:2]))

    ws2 = wb["答题明细"]
    check("答题明细行数 = 2 人 x 4 题", ws2.max_row == 9, ws2.max_row)
    head2 = [c.value for c in ws2[1]]
    expect2 = ["工号", "姓名", "题号", "题型", "题目内容", "正确答案", "考生答案", "是否正确", "得分"]
    check("答题明细表头正确", head2 == expect2, str(head2))
    d_rows = [[c.value for c in row] for row in ws2.iter_rows(min_row=2)]
    check("多选题按集合判定为正确", d_rows[3][7] == "正确" and d_rows[3][8] == 20, str(d_rows[3]))
    check("判断题答对判定正确", d_rows[2][7] == "正确" and d_rows[2][8] == 5, str(d_rows[2]))
    check("判断题答错判定错误", d_rows[6][7] == "错误" and d_rows[6][8] == 0, str(d_rows[6]))

    ws3 = wb["统计表"]
    check("统计表行数 >= 4 题 + 汇总", ws3.max_row >= 6, ws3.max_row)
    s_rows = [[c.value for c in row] for row in ws3.iter_rows(min_row=2, max_row=5)]
    check("第1题作答2人答对1人", s_rows[0][4] == 2 and s_rows[0][5] == 1, str(s_rows[0]))
    check("第4题正确率 100%", s_rows[3][6] == "100.0%", str(s_rows[3]))

    r = client.get("/api/exams/999999/export", headers=H)
    check("导出不存在的考试返回 404", r.status_code == 404, r.status_code)

    print("")
    print("== 4. 擂台赛自动判分 ==")
    r = client.post("/api/arenas/", headers=H, json={
        "title": "自动判分测试赛",
        "settings": {"required_score": 10, "rush_wrong_penalty": 5, "required_wrong_penalty": False},
        "team_names": ["甲队", "乙队"],
        "team_colors": ["#F56C6C", "#409EFF"],
    })
    check("创建擂台赛", r.status_code == 200, r.text[:300])
    aid = r.json()["id"]

    r = client.post(f"/api/arenas/{aid}/questions", headers=H, json=[
        {"question_id": q1, "q_type": "required", "order_num": 0, "score": 10},
        {"question_id": q2, "q_type": "rush", "order_num": 1, "score": 20},
        {"question_id": q3, "q_type": "required", "order_num": 2, "score": 10},
    ])
    check("添加题目", r.status_code == 200 and len(r.json()) == 3, r.text[:300])
    aq1, aq2, aq3 = [q["id"] for q in r.json()]

    teams = client.get(f"/api/arenas/{aid}/teams").json()
    t1, t2 = teams[0]["id"], teams[1]["id"]

    p1_token = client.post("/api/arenas/players/join", json={
        "team_id": t1, "device_id": "dev-nf-1", "player_name": "选手甲"}).json()["token"]
    p2_token = client.post("/api/arenas/players/join", json={
        "team_id": t2, "device_id": "dev-nf-2", "player_name": "选手乙"}).json()["token"]
    check("两队加入成功", bool(p1_token) and bool(p2_token))

    a1 = a2 = aid_q = None
    with client.websocket_connect(f"/ws/arenas/{aid}?role=host&token={token}") as host:
        with client.websocket_connect(f"/ws/arenas/{aid}?role=player&token={p1_token}") as p1:
            with client.websocket_connect(f"/ws/arenas/{aid}?role=player&token={p2_token}") as p2:
                host.send_json({"type": "show_question", "payload": {"question_id": aq1}})
                drain(host)
                p1.send_json({"type": "answer_submit", "payload": {
                    "question_id": aq1, "answer": "A", "time_used": 5}})
                p2.send_json({"type": "answer_submit", "payload": {
                    "question_id": aq1, "answer": "B", "time_used": 5}})
                sub1, _ = recv_until(host, {"answer_submitted"})
                sub2, _ = recv_until(host, {"answer_submitted"})
                check("主持人收到两条作答", sub1 is not None and sub2 is not None,
                      f"{bool(sub1)}/{bool(sub2)}")
                a1 = sub1["payload"]["answer_id"] if sub1 else None
                a2 = sub2["payload"]["answer_id"] if sub2 else None

                host.send_json({"type": "host_control", "payload": {
                    "action": "auto_judge_answer", "answer_id": a1}})
                jm, _ = recv_until(host, {"answer_judged", "auto_judge_skipped"})
                check("自动判定单条产生 answer_judged",
                      jm is not None and jm["type"] == "answer_judged",
                      str(jm["type"] if jm else None))
                if jm and jm["type"] == "answer_judged":
                    check("答对自动判为正确", jm["payload"]["is_correct"] is True,
                          str(jm["payload"]))
                    check("答对得满分 10", jm["payload"]["score"] == 10,
                          str(jm["payload"]))
                else:
                    check("答对自动判为正确", False, "无 answer_judged")
                    check("答对得满分 10", False, "无 answer_judged")

                host.send_json({"type": "host_control", "payload": {
                    "action": "auto_judge_all", "question_id": aq1}})
                done_msg, got = recv_until(host, {"auto_judge_done"})
                check("收到 auto_judge_done", done_msg is not None,
                      str([m["type"] for m in got]))
                if done_msg:
                    p = done_msg["payload"]
                    check("批量判定计数正确", p["judged_count"] == 1 and len(p["skipped"]) == 0, str(p))
                judged2 = [m for m in got if m["type"] == "answer_judged"]
                if judged2:
                    check("批量判定答错为错误", judged2[0]["payload"]["is_correct"] is False,
                          str(judged2[0]["payload"]))
                    check("必答题答错不扣分", judged2[0]["payload"]["score"] == 0,
                          str(judged2[0]["payload"]))
                else:
                    check("批量判定答错为错误", False, "无 answer_judged")
                    check("必答题答错不扣分", False, "无 answer_judged")

    db = SessionLocal()
    try:
        ans = db.query(models.TeamAnswer).filter(models.TeamAnswer.id == a1).first()
        check("数据库中已结算自动判定", ans is not None and ans.is_correct is True,
              str(ans.is_correct if ans else None))
        team = db.query(models.ArenaTeam).filter(models.ArenaTeam.id == t1).first()
        check("队伍总分同步加分", team is not None and team.score == 10,
              str(team.score if team else None))
    finally:
        db.close()

    print("")
    print("== 5. 判断题自动判分（选项无文本 -> 需人工） ==")
    r = client.post(f"/api/arenas/{aid}/questions", headers=H, json=[
        {"question_id": q5, "q_type": "required", "order_num": 3, "score": 10},
    ])
    aq4 = [q for q in r.json() if q["question_id"] == q5][0]["id"]
    with client.websocket_connect(f"/ws/arenas/{aid}?role=host&token={token}") as host:
        with client.websocket_connect(f"/ws/arenas/{aid}?role=player&token={p1_token}") as p1:
            host.send_json({"type": "show_question", "payload": {"question_id": aq4}})
            drain(host)
            p1.send_json({"type": "answer_submit", "payload": {
                "question_id": aq4, "answer": "A", "time_used": 3}})
            sub_msg, _ = recv_until(host, {"answer_submitted"})
            check("收到判断题作答", sub_msg is not None, "未收到 answer_submitted")
            if sub_msg:
                aid_q = sub_msg["payload"]["answer_id"]
                host.send_json({"type": "host_control", "payload": {
                    "action": "auto_judge_answer", "answer_id": aid_q}})
                sk_msg, got = recv_until(host, {"auto_judge_skipped", "answer_judged"})
                check("无法自动判定时提示人工评分",
                      sk_msg is not None and sk_msg["type"] == "auto_judge_skipped",
                      str(sk_msg["type"] if sk_msg else [m["type"] for m in got]))
                check("未产生 answer_judged",
                      not [m for m in got if m["type"] == "answer_judged"], "")

    if aid_q:
        db = SessionLocal()
        try:
            ans = db.query(models.TeamAnswer).filter(models.TeamAnswer.id == aid_q).first()
            check("待判定状态保持不变", ans is not None and ans.is_correct is None,
                  str(ans.is_correct if ans else None))
        finally:
            db.close()

    print("")
    print("== 6. HTTP 判分接口 is_correct=None 走自动判分 ==")
    r = client.post(f"/api/arenas/answers/{a2}/judge", headers=H, json={"is_correct": None})
    check("is_correct=None 自动判定返回 200", r.status_code == 200, r.text[:300])
    if r.status_code == 200:
        check("自动判定结果为错误", r.json().get("is_correct") is False, r.text[:200])

    print("")
    print("=" * 50)
    print(f"通过 {len(PASSED)} 项，失败 {len(FAILED)} 项")
    if FAILED:
        for f in FAILED:
            print("  [FAIL] " + f)
        sys.exit(1)
    print("全部通过")


if __name__ == "__main__":
    main()
