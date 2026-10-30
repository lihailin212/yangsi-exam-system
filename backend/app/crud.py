import json
import random
import re
import unicodedata
import string
import time
from datetime import datetime, timedelta
from typing import List
from sqlalchemy.orm import Session
from sqlalchemy import and_, func, literal_column, or_, select, union_all
from sqlalchemy.exc import IntegrityError
from . import models, schemas
from . import paper as paper_lib
from .auth import hash_password, verify_password, invalidate_user_cache
from .timeutil import (
    SUBMIT_GRACE_SECONDS,
    app_to_utc,
    to_epoch_ms,
    utc_now,
    utc_to_app,
)

# ---- 考试题目缓存 ----
#
# 题库在远端，一份 50 题的卷子传回来要一秒多，而每位考生进考试都要取一次。
# 考试期间题目是冻结的，缓存几十秒没有风险；改题、改卷时会主动作废缓存。
EXAM_QUESTION_CACHE_TTL = 30.0
_exam_question_cache = {}

# ---- User ----
def get_user_by_employee_id(db: Session, employee_id: str):
    return db.query(models.User).filter(models.User.employee_id == employee_id).first()

def get_users(db: Session, skip=0, limit=100, search="", dept=""):
    q = db.query(models.User)
    if search:
        q = q.filter(
            (models.User.name.contains(search)) | (models.User.employee_id.contains(search))
        )
    if dept:
        q = q.filter(models.User.dept == dept)
    return q.offset(skip).limit(limit).all(), q.count()

def create_user(db: Session, user_in: schemas.UserCreate):
    user = models.User(
        employee_id=user_in.employee_id,
        name=user_in.name,
        password_hash=hash_password(user_in.password),
        dept=user_in.dept,
        role=user_in.role,
        is_admin=user_in.is_admin,
        is_active=user_in.is_active,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user

def update_user(db: Session, user_id: int, user_in: schemas.UserUpdate):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user:
        return None
    if user_in.name is not None:
        user.name = user_in.name
    if user_in.dept is not None:
        user.dept = user_in.dept
    if user_in.role is not None:
        user.role = user_in.role
    if user_in.password is not None:
        user.password_hash = hash_password(user_in.password)
    if user_in.is_admin is not None:
        user.is_admin = user_in.is_admin
    if user_in.is_active is not None:
        user.is_active = user_in.is_active
    db.commit()
    db.refresh(user)
    # 管理员改了权限或启停账号，缓存里的旧快照必须立刻作废
    invalidate_user_cache(user.id)
    return user

def delete_user(db: Session, user_id: int):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if user:
        db.delete(user)
        db.commit()
        invalidate_user_cache(user.id)
    return user

# ---- Question ----
def normalize_fill_text(text: str) -> str:
    """填空题答案归一化：全角转半角、去所有空白、英文转小写。

    考生写"单 核 细 胞"、"单核细胞 "、"单核细胞" 都算对；
    全角/半角括号、英文大小写也视为同一个写法。
    """
    if not text:
        return ""
    s = unicodedata.normalize("NFKC", str(text))
    s = re.sub(r"\s+", "", s)
    # 去掉首尾的括号、引号和标点：考生写"（腺癌）"、"腺癌。"都算答对
    s = s.strip("()[]{}<>\u3010\u3011\u3008\u3009\u300a\u300b\u300c\u300d\u300e\u300f"
                "\"\u2018\u2019\u201c\u201d`\u3001\u3002\uff0c\uff0e\uff1b\uff1a\uff1f"
                "\u00b7\u2014\u2013-_+~")
    return s.strip().lower()


def fill_answer_matches(given: str, standard: str) -> bool:
    """填空题判分。标准答案可用 | 分隔多个可接受写法。"""
    g = normalize_fill_text(given)
    if not g:
        return False
    for alt in str(standard or "").split("|"):
        if g == normalize_fill_text(alt):
            return True
    return False


def get_questions(db: Session, skip=0, limit=100, q_type="", search="", category=""):
    q = db.query(models.Question)
    # 过滤掉系统分类占位题目（内容以[SYSTEM_CATEGORY]开头）
    q = q.filter(~models.Question.content.startswith("<p>[SYSTEM_CATEGORY]"))
    if q_type:
        q = q.filter(models.Question.type == q_type)
    if search:
        q = q.filter(models.Question.content.contains(search))
    if category:
        # 检查是否为父级分类（查询所有以此开头的子分类）
        if category.endswith('/'):
            # 精确匹配分类，如 "形态学/"
            q = q.filter(models.Question.category.like(f"{category}%"))
        else:
            # 精确匹配或以 "父分类/" 开头的子分类
            q = q.filter(
                (models.Question.category == category) |
                (models.Question.category.like(f"{category}/%"))
            )
    total = q.count()
    return q.order_by(models.Question.id.desc()).offset(skip).limit(limit).all(), total

def get_question(db: Session, question_id: int):
    return db.query(models.Question).filter(models.Question.id == question_id).first()

def create_question(db: Session, q_in: schemas.QuestionCreate):
    q = models.Question(**q_in.model_dump())
    db.add(q)
    db.commit()
    db.refresh(q)
    invalidate_exam_questions()
    return q

def update_question(db: Session, question_id: int, q_in: schemas.QuestionUpdate):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q:
        return None
    for k, v in q_in.model_dump().items():
        setattr(q, k, v)
    db.commit()
    db.refresh(q)
    invalidate_exam_questions()
    return q

def delete_question(db: Session, question_id: int):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if q:
        # 防止删除系统分类占位题目
        if q.content.startswith("<p>[SYSTEM_CATEGORY]"):
            return None
        db.delete(q)
        db.commit()
        invalidate_exam_questions()
    return q

def bulk_create_questions(db: Session, questions: list):
    objs = [models.Question(**q) for q in questions]
    db.add_all(objs)
    db.commit()
    invalidate_exam_questions()
    return len(objs)


def get_exam_questions(exam):
    """取考试题目，带短时缓存。exam 必须已经加载。"""
    key = exam.id
    item = _exam_question_cache.get(key)
    if item is not None:
        expires_at, questions = item
        if expires_at > time.time():
            return questions
        _exam_question_cache.pop(key, None)
    questions = list(exam.questions)
    _exam_question_cache[key] = (time.time() + EXAM_QUESTION_CACHE_TTL, questions)
    return questions


def invalidate_exam_questions(exam_id=None):
    """题目或试卷有改动时作废缓存；不传 exam_id 表示全部作废。"""
    if exam_id is None:
        _exam_question_cache.clear()
    else:
        _exam_question_cache.pop(exam_id, None)


def get_all_categories(db: Session):
    """获取所有分类及其题目数量"""
    result = db.query(
        models.Question.category,
        func.count(models.Question.id).label('count')
    ).group_by(models.Question.category).all()
    return [{"name": r.category, "count": r.count} for r in result if r.category]


def rename_category(db: Session, old_name: str, new_name: str):
    """重命名分类"""
    # 更新所有使用该分类的题目（包括系统题目）
    db.query(models.Question).filter(
        models.Question.category == old_name
    ).update({"category": new_name})
    db.commit()
    # 如果新名称为空（删除分类），删除该分类的系统占位题目
    if not new_name:
        db.query(models.Question).filter(
            models.Question.category == "",
            models.Question.content.startswith("<p>[SYSTEM_CATEGORY]")
        ).delete(synchronize_session=False)
        db.commit()


def create_category(db: Session, category_name: str):
    """创建新分类（通过创建一个系统占位题目）"""
    # 检查分类是否已存在
    existing = db.query(models.Question).filter(
        models.Question.category == category_name
    ).first()
    if existing:
        return False
    # 创建一个系统占位题目来代表这个分类（不会被普通界面显示）
    q = models.Question(
        type="single",
        content=f"<p>[SYSTEM_CATEGORY]{category_name}</p>",
        options='{"A": "系统分类"}',
        answer="A",
        analysis="系统分类占位题目，不在题目列表中显示",
        category=category_name
    )
    db.add(q)
    db.commit()
    return True


# ---- Exam ----
def get_exams(db: Session, skip=0, limit=100):
    exams = db.query(models.Exam).order_by(models.Exam.id.desc()).offset(skip).limit(limit).all()
    total = db.query(models.Exam).count()
    return exams, total

def get_exam(db: Session, exam_id: int):
    return db.query(models.Exam).filter(models.Exam.id == exam_id).first()


def _calc_total_score(exam: models.Exam) -> float:
    """计算考试总分"""
    score_map = {
        "single": exam.single_score or 10,
        "multiple": exam.multiple_score or 20,
        "judgment": exam.judgment_score or 5,
        "fill": exam.single_score or 10,
    }
    return sum(score_map.get(q.type, 10) for q in exam.questions)

def _apply_exam_questions(db: Session, exam, exam_in):
    """按提交的题目/抽题规则装配考试。

    普通考试：题目就是管理员勾选的那些。
    随机卷：题目是"候选题库"（所有规则能抽到的题），
    考生进考试时再从候选库里随机抽，题目数量按规则固定。
    """
    if exam_in.random_mode:
        rules = paper_lib.parse_rules(paper_lib.rules_to_json(exam_in.random_rules))
        if not rules:
            raise ValueError("随机卷至少要配置一条抽题规则")
        exam.random_mode = True
        exam.random_rules = paper_lib.rules_to_json(exam_in.random_rules)
        # 每条规则都要抽得出题，否则考生进考试时才发现抽不出题就晚了
        for rule in rules:
            pool = paper_lib.pool_questions(db, rule)
            if len(pool) < rule["count"]:
                raise ValueError(
                    "规则「%s」题库不足，需要 %d 道，实际只有 %d 道" % (
                        paper_lib.describe_rule(rule), rule["count"], len(pool)))
        pool_ids = paper_lib.pool_ids(db, rules)
        if not pool_ids:
            raise ValueError("抽题规则没有匹配到任何题目")
        exam.questions = db.query(models.Question).filter(
            models.Question.id.in_(pool_ids)).all()
        exam.total_score = paper_lib.estimate_total_score(db, exam, rules)
        return rules

    exam.random_mode = False
    exam.random_rules = "[]"
    if exam_in.question_ids is not None:
        qs = db.query(models.Question).filter(
            models.Question.id.in_(exam_in.question_ids)).all()
        exam.questions = qs
    exam.total_score = _calc_total_score(exam)
    return []


def create_exam(db: Session, exam_in: schemas.ExamCreate, creator_id: int):
    data = exam_in.model_dump(exclude={"question_ids", "random_rules"})
    exam = models.Exam(**data, created_by=creator_id)
    db.add(exam)
    db.flush()
    _apply_exam_questions(db, exam, exam_in)
    db.commit()
    db.refresh(exam)
    invalidate_exam_questions(exam.id)
    return exam

def update_exam(db: Session, exam_id: int, exam_in: schemas.ExamUpdate):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        return None
    data = exam_in.model_dump(exclude_none=True,
                              exclude={"question_ids", "random_rules"})
    for k, v in data.items():
        setattr(exam, k, v)
    if exam_in.random_mode is not None:
        exam.random_mode = exam_in.random_mode
    # 随机卷的规则/题目用统一入口装配，普通考试保持原来"勾选题目"的逻辑
    if exam.random_mode or exam_in.random_rules is not None:
        rules_payload = exam_in.random_rules
        if rules_payload is None and exam.random_mode:
            # 只改了标题之类、没重传规则时，沿用原来存的规则
            try:
                rules_payload = json.loads(exam.random_rules or "[]")
            except (TypeError, ValueError):
                rules_payload = []
        payload = schemas.ExamCreate(
            title=exam.title,
            duration=exam.duration,
            start_time=exam.start_time,
            end_time=exam.end_time,
            pass_score=exam.pass_score,
            single_score=exam.single_score,
            multiple_score=exam.multiple_score,
            judgment_score=exam.judgment_score,
            retake_limit=exam.retake_limit,
            random_mode=bool(exam.random_mode),
            random_rules=rules_payload or [],
            question_ids=exam_in.question_ids or [],
        )
        _apply_exam_questions(db, exam, payload)
    else:
        if exam_in.question_ids is not None:
            qs = db.query(models.Question).filter(
                models.Question.id.in_(exam_in.question_ids)).all()
            exam.questions = qs
        # 重新计算总分
        exam.total_score = _calc_total_score(exam)
    db.commit()
    db.refresh(exam)
    invalidate_exam_questions(exam_id)
    return exam

def delete_exam(db: Session, exam_id: int):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if exam:
        db.delete(exam)
        db.commit()
    invalidate_exam_questions(exam_id)
    return exam

def create_record(db: Session, exam, user_id: int, answers_str: str, at=None,
                  attempt_no=None):
    """按答案判分并生成一条交卷记录。

    at 是这条记录的提交时刻（UTC 无时区）。超时自动交卷时传截止时间，
    让成绩上的提交时间与考试截止一致，而不是结算时的时刻。
    """
    try:
        answers = json.loads(answers_str or "{}")
    except (TypeError, ValueError):
        answers = {}
    if not isinstance(answers, dict):
        answers = {}
    earned_score = 0.0

    # 根据题型映射分值
    score_map = {
        "single": exam.single_score,
        "multiple": exam.multiple_score,
        "judgment": exam.judgment_score,
        "fill": exam.single_score,
    }

    # 随机卷按该考生实际抽到的卷子判分，总分也按这份卷子算
    if getattr(exam, "random_mode", False):
        scoring = get_paper_scoring_questions(db, exam, user_id, attempt_no)
        total_score = sum(score_map.get(t, 10) for _, t, _ in scoring)
    else:
        scoring = get_scoring_questions(db, exam.id)
        total_score = exam.total_score

    for qid, qtype, correct_answer in scoring:
        question_score = score_map.get(qtype, 10)  # 默认10分
        user_ans = answers.get(str(qid), "")
        if qtype == "multiple":
            correct = set(correct_answer.split(","))
            given = set(user_ans.split(",")) if user_ans else set()
            if correct == given:
                earned_score += question_score
        elif qtype == "fill":
            if fill_answer_matches(user_ans, correct_answer):
                earned_score += question_score
        else:
            if str(user_ans).strip() == str(correct_answer).strip():
                earned_score += question_score
    record = models.ExamRecord(
        exam_id=exam.id,
        user_id=user_id,
        answers=answers_str,
        score=earned_score,
        total_score=total_score,  # 随机卷按本人卷子算，普通考试用存储的总分
    )
    if at is not None:
        record.submitted_at = at
    db.add(record)
    db.commit()
    return record


# ---- 随机卷（每人题数相同、题目随机） ----

def current_attempt_no(db: Session, exam_id: int, user_id: int) -> int:
    """该考生当前所处的作答轮次：还没开始过就是第 1 轮。"""
    latest = get_latest_attempt(db, exam_id, user_id)
    return latest.attempt_no if latest else 1


def get_or_create_paper(db: Session, exam, user_id: int, attempt_no: int):
    """取该考生该轮抽到的题目 id 列表，没有就按规则抽一份并落库。

    落库后同一个人刷新/重进都拿到同一份卷子，不同人抽到的题不同。
    """
    if not getattr(exam, "random_mode", False):
        return [q.id for q in exam.questions]

    paper = db.query(models.ExamPaper).filter(
        models.ExamPaper.exam_id == exam.id,
        models.ExamPaper.user_id == user_id,
        models.ExamPaper.attempt_no == attempt_no,
    ).first()
    if paper:
        try:
            ids = json.loads(paper.question_ids or "[]")
        except (TypeError, ValueError):
            ids = []
        if ids:
            return ids

    rules = paper_lib.parse_rules(exam.random_rules)
    ids = paper_lib.build_paper(db, rules)
    db.add(models.ExamPaper(
        exam_id=exam.id,
        user_id=user_id,
        attempt_no=attempt_no,
        question_ids=json.dumps(ids),
    ))
    try:
        db.commit()
    except IntegrityError:
        # 并发进入考试时可能同时抽卷，唯一约束保证只会留下一份
        db.rollback()
        paper = db.query(models.ExamPaper).filter(
            models.ExamPaper.exam_id == exam.id,
            models.ExamPaper.user_id == user_id,
            models.ExamPaper.attempt_no == attempt_no,
        ).first()
        if paper:
            try:
                return json.loads(paper.question_ids or "[]")
            except (TypeError, ValueError):
                pass
        raise
    return ids


def get_paper_questions(db: Session, exam, user_id: int, attempt_no=None):
    """按考生抽到的卷子取题目（保持抽题顺序）；普通考试就是考试本身的题目。"""
    if not getattr(exam, "random_mode", False):
        return list(exam.questions)
    if attempt_no is None:
        attempt_no = current_attempt_no(db, exam.id, user_id)
    ids = get_or_create_paper(db, exam, user_id, attempt_no)
    if not ids:
        return []
    rows = db.query(models.Question).filter(models.Question.id.in_(ids)).all()
    by_id = {q.id: q for q in rows}
    return [by_id[i] for i in ids if i in by_id]


def get_paper_scoring_questions(db: Session, exam, user_id: int, attempt_no=None):
    """判分用：只取该考生卷子上的题号、题型、正确答案（保持卷子顺序）。"""
    if not getattr(exam, "random_mode", False):
        return get_scoring_questions(db, exam.id)
    if attempt_no is None:
        attempt_no = current_attempt_no(db, exam.id, user_id)
    ids = get_or_create_paper(db, exam, user_id, attempt_no)
    if not ids:
        return []
    rows = db.query(
        models.Question.id,
        models.Question.type,
        models.Question.answer,
    ).filter(models.Question.id.in_(ids)).all()
    order = {qid: i for i, qid in enumerate(ids)}
    return sorted(rows, key=lambda r: order.get(r.id, 0))


def get_scoring_questions(db: Session, exam_id: int):
    """只取判分用得上的字段：题号、题型、正确答案。

    题干和选项对判分毫无用处，可一张卷子十几 KB，全量取回来要多花
    好几百毫秒——数据库在远端，载荷越大往返越慢。
    """
    return db.query(
        models.Question.id,
        models.Question.type,
        models.Question.answer,
    ).join(
        models.exam_question,
        models.exam_question.c.question_id == models.Question.id,
    ).filter(
        models.exam_question.c.exam_id == exam_id,
    ).all()


def submit_attempt(db: Session, exam, user_id: int, answers_str: str, attempt, at=None):
    """抢占作答会话并生成成绩记录。

    返回 (record, claimed)：
    - claimed 为 True：本次请求抢到了会话，record 是本次的成绩；
    - claimed 为 False：会话已被其它请求结算（并发提交或超时自动交卷），
      record 为 None，调用方应返回已有的那条成绩。

    抢占（条件 UPDATE）和落记录在同一个事务里：并发提交时只有一个请求
    能改到会话行，另一个请求会在对方提交后才返回，从而拿到同一条成绩，
    不会生成两份成绩。
    """
    ts = at or utc_now()
    changed = db.query(models.ExamAttempt).filter(
        models.ExamAttempt.id == attempt.id,
        models.ExamAttempt.submitted_at.is_(None),
    # 用 evaluate 同步会话对象：抢占成功后 attempt.submitted_at 要立刻变成
    # 已提交，否则调用方还会以为这一轮开着，白白再开一轮。
    ).update({"submitted_at": ts}, synchronize_session="evaluate")
    if not changed:
        return None, False
    return create_record(db, exam, user_id, answers_str, at=ts,
                         attempt_no=attempt.attempt_no), True


# ---- 考试会话（ExamAttempt） ----
def get_taken_count(db: Session, exam_id: int, user_id: int) -> int:
    """该考生已交卷次数。"""
    return db.query(models.ExamRecord).filter(
        models.ExamRecord.exam_id == exam_id,
        models.ExamRecord.user_id == user_id,
    ).count()


def is_exam_allowed(db: Session, exam_id: int, user_id: int) -> bool:
    """允许名单非空时，只有名单内的考生能作答。

    一次查询同时问出"名单里有没有人"和"名单里有没有这个人"，
    不用把整张名单拉回内存再找。
    """
    row = db.query(
        func.count().label("total"),
        func.count().filter(models.ExamDistribution.user_id == user_id).label("mine"),
    ).filter(
        models.ExamDistribution.exam_id == exam_id,
    ).one()
    return row.total == 0 or row.mine > 0


def get_used_rounds(db: Session, exam_id: int, user_id: int) -> int:
    """该考生已占用的作答轮次。

    一次作答轮次 = 从进入考试到交卷（或超时）这一段。
    取"交卷记录数"和"已结束会话数"的较大值：

    - 新数据里每个结束的会话都会留下一条交卷记录，两者相等；
    - 历史数据只有交卷记录、没有会话，用记录数兜底；
    - 超时的会话也算用掉一轮，轮次限制才不会漏。
    """
    used, _ = get_exam_progress(db, exam_id, user_id)
    return used


def get_answered_count(db: Session, exam_id: int, user_id: int) -> int:
    """该考生真正作答过的交卷次数（答案非空的记录）。

    与"已用轮次"分开：超时后一道题都没答的考生占用了一轮，
    但不能因此就让他看到正确答案。
    """
    _, answered = get_exam_progress(db, exam_id, user_id)
    return answered


def get_exam_progress(db: Session, exam_id: int, user_id: int):
    """一次数据库往返算出"已用轮次"和"已作答轮次"。

    这两个数原本要查三次（交卷记录数、已结束会话数、逐条读答案），
    而数据库在远端，一次往返两百多毫秒，三次就是近一秒。合成一条
    SQL 后只剩一次往返：

    - 前半段每行一条交卷记录的答案，用来数出答案非空的记录数；
    - 后半段是一行聚合，永远存在，给出已结束会话数。

    返回 (used_rounds, answered_count)。
    """
    branch_records = select(
        models.ExamRecord.answers.label("answers"),
        literal_column("0").label("kind"),
        literal_column("NULL::bigint").label("cnt"),
    ).select_from(models.ExamRecord).where(
        models.ExamRecord.exam_id == exam_id,
        models.ExamRecord.user_id == user_id,
    )
    branch_closed = select(
        literal_column("NULL::text").label("answers"),
        literal_column("1").label("kind"),
        func.count().label("cnt"),
    ).select_from(models.ExamAttempt).where(
        models.ExamAttempt.exam_id == exam_id,
        models.ExamAttempt.user_id == user_id,
        models.ExamAttempt.submitted_at.isnot(None),
    )
    rows = db.execute(union_all(branch_records, branch_closed)).all()

    records = 0
    answered = 0
    closed = 0
    for row in rows:
        if row.kind == 1:
            closed = int(row.cnt or 0)
            continue
        records += 1
        raw = row.answers
        if isinstance(raw, (bytes, bytearray)):
            raw = raw.decode("utf-8", "replace")
        try:
            parsed = json.loads(raw or "{}")
        except (TypeError, ValueError):
            parsed = {}
        if isinstance(parsed, dict) and parsed:
            answered += 1
    # 已用轮次取两者较大值：新数据里结束的会话都留下一条交卷记录，
    # 两者相等；历史数据只有交卷记录，用记录数兜底。
    return max(records, closed), answered


def get_latest_attempt(db: Session, exam_id: int, user_id: int):
    """取该考生最近一次作答会话（不限是否已交卷）。"""
    return db.query(models.ExamAttempt).filter(
        models.ExamAttempt.exam_id == exam_id,
        models.ExamAttempt.user_id == user_id,
    ).order_by(models.ExamAttempt.attempt_no.desc()).first()


def get_open_attempt(db: Session, exam_id: int, user_id: int):
    """取当前这一轮尚未结束的会话（可能已经超时）。"""
    latest = get_latest_attempt(db, exam_id, user_id)
    if latest is None or latest.submitted_at is not None:
        return None
    return latest


def count_closed_attempts(db: Session, exam_id: int, user_id: int) -> int:
    """已结束的作答轮次（已交卷 + 超时未交卷）。"""
    return db.query(models.ExamAttempt).filter(
        models.ExamAttempt.exam_id == exam_id,
        models.ExamAttempt.user_id == user_id,
        models.ExamAttempt.submitted_at.isnot(None),
    ).count()


def is_attempt_expired(attempt) -> bool:
    """作答会话是否已过期（含交卷宽限期）。

    宽限期给客户端留了提交的网络延迟余量：刚过截止时间几十秒内提交
    仍算有效，超过宽限期才算超时。草稿保存、超时结算、提交三条路径
    都用这一个判定，结论必须一致。
    """
    if attempt is None or attempt.deadline_at is None:
        return False
    return (utc_now() - attempt.deadline_at).total_seconds() > SUBMIT_GRACE_SECONDS


def finalize_expired_attempt(db: Session, exam, user_id: int, attempt=None):
    """结算超时会话：按服务端已保存的作答自动交卷。

    没有任何作答也要记一条 0 分成绩，这样"进入过考试"和"已用轮次"
    才能一一对应，考生反复超时重进也占不掉补考机会以外的额外时间。
    attempt 是调用方已经取到的会话，省掉一次查询。
    返回生成的交卷记录；没有超时则返回 None。
    """
    if attempt is None:
        attempt = get_open_attempt(db, exam.id, user_id)
    if attempt is None or attempt.submitted_at is not None or not is_attempt_expired(attempt):
        return None
    # 抢占会话，抢不到说明已被其它请求结算过，避免重复生成成绩
    record, claimed = submit_attempt(
        db, exam, user_id, attempt.draft or "{}", attempt, at=attempt.deadline_at
    )
    if not claimed:
        return None
    return record


def settle_and_count_rounds(db: Session, exam, user_id: int):
    """结算超时会话后，返回 (该考生已用的作答轮次, 最近一次会话)。

    超时的会话同样占一轮。如果只在"有交卷记录"时才计数，
    考生超时后重进就会白赚一轮，总作答时长被无限延长。
    开考、考生登录、交卷三处都用这个口径，次数判断才不会前后不一。

    把最近一次会话一并返回，是为了让调用方少查一次库；
    超时被结算过的会话，它的 submitted_at 也已经是最新状态。
    """
    latest = get_latest_attempt(db, exam.id, user_id)
    if latest is not None and latest.submitted_at is None and is_attempt_expired(latest):
        finalize_expired_attempt(db, exam, user_id, attempt=latest)
    return get_used_rounds(db, exam.id, user_id), latest


def create_attempt(db: Session, exam, user_id: int, attempt_no=None):
    """开启一轮新的作答会话。

    attempt_no 是本轮的序号；不传就自己查一次最近一次会话再续。
    """
    if attempt_no is None:
        latest = get_latest_attempt(db, exam.id, user_id)
        attempt_no = (latest.attempt_no + 1) if latest else 1
    started_at = utc_now()
    deadline = started_at + timedelta(minutes=exam.duration or 0)
    # 考试总时长不得超过考试窗口的结束时间
    if exam.end_time:
        window_end = app_to_utc(exam.end_time)
        if window_end < deadline:
            deadline = window_end
    attempt = models.ExamAttempt(
        exam_id=exam.id,
        user_id=user_id,
        attempt_no=attempt_no,
        started_at=started_at,
        deadline_at=deadline,
    )
    db.add(attempt)
    db.commit()
    return attempt


def start_attempt(db: Session, exam, user_id: int, latest=None):
    """进入考试时开启（或复用）一次作答会话。

    幂等：同一轮作答重复进入不会重置开始时间，考生反复登录
    或刷新页面都无法把考试时限重新算一遍。超时的那一轮会先结掉
    （按草稿自动交卷），再开新一轮。

    exam 复用调用方已经加载的考试对象，latest 是已经取到的最近会话，
    两者都能省掉一次数据库往返。
    """
    if not exam:
        return None
    if latest is None:
        latest = get_latest_attempt(db, exam.id, user_id)
    next_no = (latest.attempt_no + 1) if latest else 1
    if latest is not None and latest.submitted_at is None and is_attempt_expired(latest):
        # 超时的那一轮先按草稿自动交卷，然后才开新一轮
        finalize_expired_attempt(db, exam, user_id, attempt=latest)
        return create_attempt(db, exam, user_id, attempt_no=next_no)
    if latest is not None and latest.submitted_at is None:
        # 未超时：复用当前会话，倒计时不重置
        return latest
    return create_attempt(db, exam, user_id, attempt_no=next_no)


def get_latest_record(db: Session, exam_id: int, user_id: int):
    """该考生最近一次交卷记录。"""
    return db.query(models.ExamRecord).filter(
        models.ExamRecord.exam_id == exam_id,
        models.ExamRecord.user_id == user_id,
    ).order_by(models.ExamRecord.id.desc()).first()

def get_user_records(db: Session, user_id: int):
    return db.query(models.ExamRecord).filter(models.ExamRecord.user_id == user_id).all()

def get_exam_records(db: Session, exam_id: int):
    return db.query(models.ExamRecord).filter(models.ExamRecord.exam_id == exam_id).all()

# ---- Statistics ----
def _percent(score, total_score):
    """得分率。满分缺失或为 0 时按 0 算，避免除零。"""
    if not total_score:
        return 0.0
    return (score or 0) / total_score * 100


def get_statistics(db: Session):
    """汇总统计。

    这里所有派生数字都从同一条 JOIN 出来的记录流里算，不再一条条回查。
    以前及格率是"每条记录再查一次考试"，近 7 天趋势是"每天查一次"，
    一次请求要打 19 次库；数据库在境外，每次往返都是几十到几百毫秒，
    结果这个接口要 7 秒，前端 15 秒超时就直接失败，页面一片 0。
    """
    ExamRecord, Exam = models.ExamRecord, models.Exam

    total_users = db.query(models.User).filter(models.User.is_admin == False).count()
    total_exams = db.query(Exam).count()

    rows = (
        db.query(ExamRecord.score, ExamRecord.total_score, Exam.pass_score)
        .outerjoin(Exam, Exam.id == ExamRecord.exam_id)
        .all()
    )
    total_records = len(rows)

    avg_score = 0.0
    pass_rate = 0.0
    score_dist = {"90-100": 0, "80-89": 0, "70-79": 0, "60-69": 0, "below60": 0}
    if total_records:
        total_pct = 0.0
        pass_count = 0
        for score, total_score, pass_score in rows:
            pct = _percent(score, total_score)
            total_pct += pct
            if pass_score is not None and pct >= pass_score:
                pass_count += 1
            if pct >= 90:
                score_dist["90-100"] += 1
            elif pct >= 80:
                score_dist["80-89"] += 1
            elif pct >= 70:
                score_dist["70-79"] += 1
            elif pct >= 60:
                score_dist["60-69"] += 1
            else:
                score_dist["below60"] += 1
        avg_score = total_pct / total_records
        pass_rate = pass_count / total_records * 100

    # 近 7 天趋势：一条 GROUP BY 查完，再补上没有交卷的 0
    today = datetime.utcnow().date()
    days = [today - timedelta(days=i) for i in range(6, -1, -1)]
    grouped = (
        db.query(
            func.date(ExamRecord.submitted_at).label("day"),
            func.count(ExamRecord.id),
        )
        .filter(func.date(ExamRecord.submitted_at) >= days[0])
        .group_by("day")
        .all()
    )
    by_day = {r.day: r[1] for r in grouped}
    trend = [{"date": str(d), "count": by_day.get(d, 0)} for d in days]

    now = utc_now()
    cur_start = now - timedelta(days=7)
    prev_start = now - timedelta(days=14)

    # 上一个 7 天周期的及格率，用于算"及格率"卡片的百分点变化
    prev_rows = (
        db.query(ExamRecord.score, ExamRecord.total_score, Exam.pass_score)
        .outerjoin(Exam, Exam.id == ExamRecord.exam_id)
        .filter(ExamRecord.submitted_at >= prev_start, ExamRecord.submitted_at < cur_start)
        .all()
    )
    prev_pass_rate = 0.0
    if prev_rows:
        prev_pass = 0
        for score, total_score, pass_score in prev_rows:
            pct = _percent(score, total_score)
            if pass_score is not None and pct >= pass_score:
                prev_pass += 1
        prev_pass_rate = prev_pass / len(prev_rows) * 100

    # 首页卡片的涨跌箭头：本 7 天 vs 上 7 天。以前是写死的 12% / 8% / -3%，
    # 涨跌和实际数据没有任何关系，这里按真实增量算。
    def _delta_pct(cur, prev):
        # 上一个周期没有数据时算不出百分比（0 -> 3 是多少百分比？没有意义），
        # 返回 None，前端改显示"新增 N"，别拿个假的 100% 糊弄人。
        if not prev:
            return None
        return round((cur - prev) / prev * 100, 1)

    cur_users = (
        db.query(func.count(models.User.id))
        .filter(models.User.is_admin == False, models.User.created_at >= cur_start)
        .scalar()
    ) or 0
    prev_users = (
        db.query(func.count(models.User.id))
        .filter(
            models.User.is_admin == False,
            models.User.created_at >= prev_start,
            models.User.created_at < cur_start,
        )
        .scalar()
    ) or 0
    cur_exams = (
        db.query(func.count(Exam.id)).filter(Exam.created_at >= cur_start).scalar()
    ) or 0
    prev_exams = (
        db.query(func.count(Exam.id))
        .filter(Exam.created_at >= prev_start, Exam.created_at < cur_start)
        .scalar()
    ) or 0
    cur_records = (
        db.query(func.count(ExamRecord.id))
        .filter(ExamRecord.submitted_at >= cur_start)
        .scalar()
    ) or 0
    prev_records = (
        db.query(func.count(ExamRecord.id))
        .filter(ExamRecord.submitted_at >= prev_start, ExamRecord.submitted_at < cur_start)
        .scalar()
    ) or 0

    pass_delta = None if not prev_rows else round(pass_rate - prev_pass_rate, 1)
    deltas = {
        "users": _delta_pct(cur_users, prev_users),
        "exams": _delta_pct(cur_exams, prev_exams),
        "records": _delta_pct(cur_records, prev_records),
        "pass": pass_delta,
        "counts": {
            "users": cur_users,
            "exams": cur_exams,
            "records": cur_records,
        },
    }

    # 首页"学习进度"四条：以前也是写死的 78% / 92% / 65% / 85%，改成真实指标。
    participated = (
        db.query(func.count(func.distinct(ExamRecord.user_id))).scalar()
    ) or 0
    participation_rate = round(participated / total_users * 100, 1) if total_users else 0.0
    month_ago = now - timedelta(days=30)
    active_exams = (
        db.query(func.count(func.distinct(ExamRecord.exam_id)))
        .filter(ExamRecord.submitted_at >= month_ago)
        .scalar()
    ) or 0
    month_exam_rate = round(active_exams / total_exams * 100, 1) if total_exams else 0.0

    return {
        "total_users": total_users,
        "total_exams": total_exams,
        "total_records": total_records,
        "avg_score": round(avg_score, 1),
        "pass_rate": round(pass_rate, 1),
        "trend": trend,
        "score_dist": score_dist,
        "deltas": deltas,
        "progress": {
            "participation_rate": participation_rate,
            "month_exam_rate": month_exam_rate,
            "avg_score_rate": round(avg_score, 1),
            "pass_rate": round(pass_rate, 1),
        },
    }


def get_recent_activity(db: Session, limit=10):
    """首页"最近动态"：把系统里最近真发生过的几件事按时间倒序拼出来。

    之前这个列表是写死的演示数据（"张三完成了…""李四开始作答…"），
    页面看着热闹，其实一条真实行为都没有，管理员会以为系统没在记录。
    现在从三张表里取真事：

      - exam_records.submitted_at  → 交卷（完成考试）
      - exam_attempts.started_at    → 进入考试（开始作答，尚未交卷）
      - users.created_at            → 新学员注册

    时间统一按考试所在地时区（北京时间）返回，同时带一个毫秒时间戳，
    前端算"X 分钟前"用时间戳，显示用格式化后的字符串。
    """
    ExamRecord, User, Exam, Attempt = (
        models.ExamRecord, models.User, models.Exam, models.ExamAttempt
    )

    events = []

    rows = (
        db.query(
            User.name,
            User.employee_id,
            Exam.title,
            ExamRecord.score,
            ExamRecord.total_score,
            ExamRecord.submitted_at,
        )
        .join(User, User.id == ExamRecord.user_id)
        .outerjoin(Exam, Exam.id == ExamRecord.exam_id)
        .order_by(ExamRecord.submitted_at.desc(), ExamRecord.id.desc())
        .limit(limit)
        .all()
    )
    for name, employee_id, title, score, total_score, submitted_at in rows:
        events.append({
            "type": "exam_completed",
            "user_name": name or employee_id or "考生",
            "exam_title": title or "（考试已删除）",
            "score": score,
            "total_score": total_score,
            "ts": _ts_ms(submitted_at),
            "time": _fmt_app_time(submitted_at),
        })

    rows = (
        db.query(
            User.name,
            User.employee_id,
            Exam.title,
            Attempt.started_at,
        )
        .join(User, User.id == Attempt.user_id)
        .outerjoin(Exam, Exam.id == Attempt.exam_id)
        .filter(Attempt.submitted_at.is_(None))
        .order_by(Attempt.started_at.desc(), Attempt.id.desc())
        .limit(limit)
        .all()
    )
    for name, employee_id, title, started_at in rows:
        events.append({
            "type": "exam_started",
            "user_name": name or employee_id or "考生",
            "exam_title": title or "（考试已删除）",
            "score": None,
            "total_score": None,
            "ts": _ts_ms(started_at),
            "time": _fmt_app_time(started_at),
        })

    rows = (
        db.query(User.name, User.employee_id, User.created_at)
        .filter(User.is_admin == False)
        .order_by(User.created_at.desc(), User.id.desc())
        .limit(limit)
        .all()
    )
    for name, employee_id, created_at in rows:
        events.append({
            "type": "user_joined",
            "user_name": name or employee_id or "新学员",
            "exam_title": None,
            "score": None,
            "total_score": None,
            "ts": _ts_ms(created_at),
            "time": _fmt_app_time(created_at),
        })

    events.sort(key=lambda e: e["ts"], reverse=True)
    return events[:limit]


def _ts_ms(dt):
    """UTC 无时区时间 → 毫秒时间戳（和 to_epoch_ms 一个口径）。"""
    if dt is None:
        return 0
    return to_epoch_ms(dt)


def _fmt_app_time(dt):
    """UTC 无时区时间 → 北京时间字符串（页面直接显示，不用再换算）。"""
    if dt is None:
        return ""
    return utc_to_app(dt).strftime("%Y-%m-%d %H:%M:%S")


def get_statistics_trend(db: Session, days=7):
    """交卷量趋势。

    days=7/30 按天返回，days=365 按月返回（一年 365 个柱子没有意义）。
    """
    ExamRecord = models.ExamRecord
    today = utc_now().date()
    days = max(1, min(int(days or 7), 400))

    if days <= 60:
        start = today - timedelta(days=days - 1)
        grouped = (
            db.query(
                func.date(ExamRecord.submitted_at).label("day"),
                func.count(ExamRecord.id),
            )
            .filter(func.date(ExamRecord.submitted_at) >= start)
            .group_by("day")
            .all()
        )
        by_day = {str(r.day): r[1] for r in grouped}
        buckets = []
        for i in range(days):
            d = start + timedelta(days=i)
            key = d.isoformat()
            buckets.append({
                "key": key,
                "label": d.strftime("%m-%d"),
                "count": by_day.get(key, 0),
            })
        return buckets

    # 按月：近 12 个月（含本月）。
    # 注意窗口要算准：之前起点取的是"上个月 1 号"，
    # 结果本月的交卷被排除在统计外，最后一个月永远是 0。
    cur_month = today.replace(day=1)
    months = []
    for i in range(11, -1, -1):
        m = cur_month
        for _ in range(i):
            m = (m - timedelta(days=1)).replace(day=1)
        months.append(m)
    start = months[0]

    grouped = (
        db.query(
            func.to_char(ExamRecord.submitted_at, "YYYY-MM").label("month"),
            func.count(ExamRecord.id),
        )
        .filter(ExamRecord.submitted_at >= start)
        .group_by("month")
        .all()
    )
    by_month = {str(r.month): r[1] for r in grouped}
    buckets = []
    for m in months:
        key = m.strftime("%Y-%m")
        buckets.append({
            "key": key,
            "label": m.strftime("%y-%m"),
            "count": by_month.get(key, 0),
        })
    return buckets


def get_statistics_records(db: Session, exam_id=None, user_id=None, keyword=None, limit=500):
    """每一条交卷记录的明细：给数据统计页的"考生答题情况"表用。

    用时来自 exam_attempts（进入考试写 started_at、交卷写 submitted_at），
    那是服务端的时间，考生清缓存重登也改不了。作答会话机制上线前的历史
    记录没有这条数据，用时显示为空。

    答题次数直接按 (exam_id, user_id) 分组排序算，历史记录也算得出来，
    不依赖 exam_attempts。
    """
    ExamRecord, User, Exam, Attempt = models.ExamRecord, models.User, models.Exam, models.ExamAttempt

    attempt_no = func.row_number().over(
        partition_by=(ExamRecord.exam_id, ExamRecord.user_id),
        order_by=(ExamRecord.submitted_at, ExamRecord.id),
    ).label("attempt_no")

    q = (
        db.query(
            ExamRecord.id,
            User.name,
            User.employee_id,
            Exam.title,
            Exam.pass_score,
            ExamRecord.score,
            ExamRecord.total_score,
            ExamRecord.submitted_at,
            attempt_no,
            Attempt.started_at,
        )
        .join(User, User.id == ExamRecord.user_id)
        .outerjoin(Exam, Exam.id == ExamRecord.exam_id)
        # 会话和成绩记录是同一时刻落库的（submit_attempt 里同一个 ts），
        # 用 submitted_at 对齐就能取到这次作答的开始时间，算出用时。
        .outerjoin(
            Attempt,
            and_(
                Attempt.exam_id == ExamRecord.exam_id,
                Attempt.user_id == ExamRecord.user_id,
                Attempt.submitted_at == ExamRecord.submitted_at,
            ),
        )
    )

    if exam_id:
        q = q.filter(ExamRecord.exam_id == exam_id)
    if user_id:
        q = q.filter(ExamRecord.user_id == user_id)
    if keyword:
        like = f"%{keyword}%"
        q = q.filter(or_(User.name.ilike(like), User.employee_id.ilike(like), Exam.title.ilike(like)))

    rows = q.order_by(ExamRecord.submitted_at.desc(), ExamRecord.id.desc()).limit(limit).all()

    result = []
    for r in rows:
        (rid, name, employee_id, title, pass_score,
         score, total_score, submitted_at, no, started_at) = r

        used_seconds = None
        if started_at is not None and submitted_at is not None:
            used_seconds = int((submitted_at - started_at).total_seconds())
            if used_seconds < 0:
                used_seconds = None

        pct = _percent(score, total_score)
        result.append({
            "record_id": rid,
            "name": name,
            "employee_id": employee_id,
            "exam_title": title or "（考试已删除）",
            "score": round(score or 0, 1),
            "total_score": round(total_score or 0, 1),
            "percent": round(pct, 1),
            "pass_score": pass_score,
            "passed": bool(pass_score is not None and pct >= pass_score),
            "attempt_no": no,
            "used_seconds": used_seconds,
            "submitted_at": submitted_at,
        })
    return result

# ---- ExamDistribution ----
def _generate_invite_code(db: Session) -> str:
    """生成唯一的6位邀请码"""
    while True:
        code = ''.join(random.choices(string.ascii_uppercase + string.digits, k=6))
        exists = db.query(models.ExamDistribution).filter(
            models.ExamDistribution.invite_code == code
        ).first()
        if not exists:
            return code

def create_distributions(db: Session, exam_id: int, user_ids: list):
    """批量创建分发记录，同一考生参加同一考试共用一个邀请码"""
    records = []
    for user_id in user_ids:
        # 检查是否已存在该分发
        existing = db.query(models.ExamDistribution).filter(
            models.ExamDistribution.exam_id == exam_id,
            models.ExamDistribution.user_id == user_id
        ).first()
        if existing:
            records.append(existing)
            continue
        invite_code = _generate_invite_code(db)
        dist = models.ExamDistribution(
            exam_id=exam_id,
            user_id=user_id,
            invite_code=invite_code,
        )
        db.add(dist)
        records.append(dist)
    db.commit()
    for r in records:
        db.refresh(r)
    return records

def get_distributions(db: Session, exam_id: int):
    """获取某考试的所有分发记录"""
    return db.query(models.ExamDistribution).filter(
        models.ExamDistribution.exam_id == exam_id
    ).all()

def delete_distribution(db: Session, distribution_id: int):
    """删除分发记录"""
    dist = db.query(models.ExamDistribution).filter(
        models.ExamDistribution.id == distribution_id
    ).first()
    if dist:
        db.delete(dist)
        db.commit()
    return dist

def get_distribution_by_code(db: Session, code: str):
    """通过邀请码查询分发记录"""
    return db.query(models.ExamDistribution).filter(
        models.ExamDistribution.invite_code == code
    ).first()

def verify_distribution_user(db: Session, code: str, employee_id: str, password: str):
    """验证邀请码+用户凭证，返回 (user, distribution) 或 (None, None)"""
    dist = get_distribution_by_code(db, code)
    if not dist:
        return None, None
    user = get_user_by_employee_id(db, employee_id)
    if not user:
        return None, None
    if not verify_password(password, user.password_hash):
        return None, None
    # 检查该用户是否在此分发记录中
    if dist.user_id != user.id:
        return None, None
    return user, dist

# ========== 专业擂台赛 CRUD ==========

def _generate_arena_invite_code(db: Session) -> str:
    """生成唯一的6位邀请码"""
    while True:
        code = ''.join(random.choices(string.digits, k=6))  # 纯数字
        exists = db.query(models.Arena).filter(models.Arena.invite_code == code).first()
        if not exists:
            return code

def create_arena(db: Session, arena_in: schemas.ArenaCreate, creator_id: int):
    """创建擂台赛"""
    settings_dict = arena_in.settings.model_dump() if isinstance(arena_in.settings, schemas.ArenaSettings) else arena_in.settings
    arena = models.Arena(
        title=arena_in.title,
        settings=json.dumps(settings_dict),
        invite_code=_generate_arena_invite_code(db),
        invite_expire=datetime.utcnow() + timedelta(minutes=15),
        created_by=creator_id,
    )
    db.add(arena)
    db.commit()
    db.refresh(arena)

    # 创建队伍
    for i, name in enumerate(arena_in.team_names):
        color = arena_in.team_colors[i] if i < len(arena_in.team_colors) else get_default_team_color(i)
        team = models.ArenaTeam(
            arena_id=arena.id,
            name=name,
            order=i,
            color=color,
        )
        db.add(team)
    db.commit()
    db.refresh(arena)
    return arena

def get_default_team_color(index: int) -> str:
    """获取默认队伍颜色"""
    colors = ["#409EFF", "#67C23A", "#E6A23C", "#F56C6C", "#909399", "#8E44AD", "#3498DB", "#E74C3C", "#2ECC71", "#F39C12"]
    return colors[index % len(colors)]

def get_arenas(db: Session, skip=0, limit=100, status=""):
    """获取擂台赛列表"""
    q = db.query(models.Arena)
    if status:
        q = q.filter(models.Arena.status == status)
    total = q.count()
    arenas = q.order_by(models.Arena.id.desc()).offset(skip).limit(limit).all()
    return arenas, total

def get_arena(db: Session, arena_id: int):
    """获取擂台赛详情"""
    return db.query(models.Arena).filter(models.Arena.id == arena_id).first()

def update_arena(db: Session, arena_id: int, arena_in: schemas.ArenaUpdate):
    """更新擂台赛"""
    arena = get_arena(db, arena_id)
    if not arena:
        return None
    if arena_in.title is not None:
        arena.title = arena_in.title
    if arena_in.status is not None:
        arena.status = arena_in.status
    if arena_in.settings is not None:
        if isinstance(arena_in.settings, schemas.ArenaSettings):
            arena.settings = json.dumps(arena_in.settings.model_dump())
        else:
            arena.settings = json.dumps(arena_in.settings)
    db.commit()
    db.refresh(arena)
    return arena

def delete_arena(db: Session, arena_id: int):
    """删除擂台赛"""
    arena = get_arena(db, arena_id)
    if arena:
        db.delete(arena)
        db.commit()
    return arena

def get_arena_by_invite_code(db: Session, code: str):
    """通过邀请码获取擂台赛"""
    arena = db.query(models.Arena).filter(models.Arena.invite_code == code).first()
    if not arena:
        return None
    if not is_invite_valid(arena):
        return None
    return arena

def is_invite_valid(arena) -> bool:
    """邀请码有效性：15分钟有效 + 比赛结束后失效"""
    if not arena:
        return False
    if arena.status == "finished":
        return False
    if arena.invite_expire and arena.invite_expire < datetime.utcnow():
        return False
    return True

def refresh_arena_invite(db: Session, arena_id: int):
    """刷新邀请码并重置有效期"""
    arena = get_arena(db, arena_id)
    if not arena:
        return None
    arena.invite_code = _generate_arena_invite_code(db)
    arena.invite_expire = datetime.utcnow() + timedelta(minutes=15)
    db.commit()
    db.refresh(arena)
    return arena

def add_team_to_arena(db: Session, arena_id: int, team_in: schemas.ArenaTeamCreate):
    """添加队伍到擂台赛"""
    arena = get_arena(db, arena_id)
    if not arena:
        return None
    max_order = db.query(models.ArenaTeam).filter(
        models.ArenaTeam.arena_id == arena_id
    ).count()
    team = models.ArenaTeam(
        arena_id=arena_id,
        name=team_in.name,
        color=team_in.color,
        order=max_order,
    )
    db.add(team)
    db.commit()
    db.refresh(team)
    return team

def get_arena_teams(db: Session, arena_id: int):
    """获取擂台赛所有队伍"""
    return db.query(models.ArenaTeam).filter(
        models.ArenaTeam.arena_id == arena_id
    ).order_by(models.ArenaTeam.order).all()

def get_arena_team(db: Session, team_id: int):
    """获取队伍详情"""
    return db.query(models.ArenaTeam).filter(models.ArenaTeam.id == team_id).first()

def update_arena_team(db: Session, team_id: int, team_in: schemas.ArenaTeamUpdate):
    """更新队伍信息"""
    team = get_arena_team(db, team_id)
    if not team:
        return None
    if team_in.name is not None:
        team.name = team_in.name
    if team_in.color is not None:
        team.color = team_in.color
    if team_in.score is not None:
        team.score = team_in.score
    if team_in.status is not None:
        team.status = team_in.status
    if team_in.help_used_this_round is not None:
        team.help_used_this_round = team_in.help_used_this_round
    db.commit()
    db.refresh(team)
    return team

def reset_team_help_status(db: Session, arena_id: int):
    """重置所有队伍本轮求助状态"""
    db.query(models.ArenaTeam).filter(
        models.ArenaTeam.arena_id == arena_id
    ).update({"help_used_this_round": False})
    db.commit()

def add_questions_to_arena(db: Session, arena_id: int, questions: List[schemas.ArenaQuestionCreate]):
    """添加题目到擂台赛"""
    arena = get_arena(db, arena_id)
    if not arena:
        return []
    for q in questions:
        # 检查题目是否已存在
        existing = db.query(models.ArenaQuestion).filter(
            models.ArenaQuestion.arena_id == arena_id,
            models.ArenaQuestion.question_id == q.question_id,
        ).first()
        if existing:
            continue
        aq = models.ArenaQuestion(
            arena_id=arena_id,
            question_id=q.question_id,
            q_type=q.q_type,
            order_num=q.order_num,
            score=q.score,
        )
        db.add(aq)
    db.commit()
    return get_arena_questions(db, arena_id)

def get_arena_questions(db: Session, arena_id: int):
    """获取擂台赛所有题目"""
    return db.query(models.ArenaQuestion).filter(
        models.ArenaQuestion.arena_id == arena_id
    ).order_by(models.ArenaQuestion.order_num).all()

def get_arena_question(db: Session, question_id: int):
    """获取擂台赛题目"""
    return db.query(models.ArenaQuestion).filter(models.ArenaQuestion.id == question_id).first()

def submit_team_answer(db: Session, answer_in: schemas.TeamAnswerSubmit):
    """提交队伍答案（只记录，不判分——判分由主持人完成）"""
    team = get_arena_team(db, answer_in.team_id)
    if not team:
        return None
    arena_question = get_arena_question(db, answer_in.question_id)
    if not arena_question:
        return None
    # 题目必须属于该队所在擂台赛
    if arena_question.arena_id != team.arena_id:
        return None
    # 检查是否已回答
    existing = db.query(models.TeamAnswer).filter(
        models.TeamAnswer.team_id == answer_in.team_id,
        models.TeamAnswer.question_id == answer_in.question_id,
    ).first()
    if existing:
        return existing
    answer = models.TeamAnswer(
        team_id=answer_in.team_id,
        question_id=answer_in.question_id,
        answer=answer_in.answer,
        is_correct=None,
        time_used=answer_in.time_used,
        score=0,
        help_used=answer_in.help_used or "",
    )
    db.add(answer)
    db.commit()
    db.refresh(answer)
    return answer

def get_team_answers(db: Session, team_id: int):
    """获取队伍所有答题记录"""
    return db.query(models.TeamAnswer).filter(
        models.TeamAnswer.team_id == team_id
    ).all()


def get_pending_answers(db: Session, arena_id: int, question_id: int):
    """获取某题目下已提交但尚未判定的作答（主持人断线重连后恢复待判列表）"""
    return db.query(models.TeamAnswer).join(
        models.ArenaQuestion, models.TeamAnswer.question_id == models.ArenaQuestion.id
    ).filter(
        models.ArenaQuestion.arena_id == arena_id,
        models.TeamAnswer.question_id == question_id,
        models.TeamAnswer.is_correct.is_(None),
    ).order_by(models.TeamAnswer.id).all()

def get_arena_settings(arena) -> dict:
    """解析擂台赛规则配置"""
    if not arena or not arena.settings:
        return {}
    try:
        data = json.loads(arena.settings)
    except (json.JSONDecodeError, TypeError):
        return {}
    return data if isinstance(data, dict) else {}

def auto_judge_answer(arena_question, original_question, answer: str):
    """自动比对客观题答案；无法自动判定（主观题/无标准答案/选项无文本）时返回 None 交给主持人"""
    if not original_question or not original_question.answer:
        return None
    if arena_question.q_type == "analysis":
        return None
    return _auto_judge_objective(original_question, answer)


def _resolve_option_text(question, value: str) -> str:
    """把选项键或选项文本统一解析成"选项文本"，便于与标准答案比对。"""
    raw = (value or "").strip()
    if not raw:
        return ""
    try:
        options = json.loads(question.options or "{}")
    except (json.JSONDecodeError, TypeError):
        options = {}
    if isinstance(options, dict):
        for key, text in options.items():
            if str(key).strip().upper() == raw.upper():
                return str(text or "").strip()
    return raw


def _auto_judge_objective(question, given_answer: str):
    """客观题自动比对，返回 True / False / None（无法判定）。

    None 表示题目形态不适合机器判定（例如判断题选项没有文本，无法确定
    选手选的字母对应"正确"还是"错误"），此时必须由主持人人工评分。
    """
    standard = (question.answer or "").strip()
    if not standard:
        return None
    given = (given_answer or "").strip()

    if question.type == "multiple":
        correct_set = {x.strip().upper() for x in standard.split(",") if x.strip()}
        given_set = {x.strip().upper() for x in given.split(",") if x.strip()}
        return correct_set == given_set

    if question.type == "fill":
        return fill_answer_matches(given, standard)

    if not given:
        return False

    # 选手直接答了选项文本（如判断题答"正确"）
    if given.upper() == standard.upper():
        return True

    # 选手答了选项键（如"A"），标准答案可能是键也可能是文本
    standard_text = _resolve_option_text(question, standard)
    given_text = _resolve_option_text(question, given)
    if not given_text:
        # 选项没有文本（例如判断题选项为空），无法确定含义
        return None
    if given_text.upper() == standard_text.upper():
        return True
    return False


def auto_judge_pending(db: Session, arena_id: int, question_id: int):
    """批量自动判定某题下所有待判定的作答。

    返回 (judged, skipped)：judged 为已自动判定成功的作答列表，
    skipped 为 [{"answer_id":..., "team_name":..., "reason":...}]（需人工判定）。
    """
    teams = {t.id: t for t in get_arena_teams(db, arena_id)}
    judged = []
    skipped = []
    for ans in get_pending_answers(db, arena_id, question_id):
        arena_question = get_arena_question(db, ans.question_id)
        if not arena_question:
            continue
        original = db.query(models.Question).filter(
            models.Question.id == arena_question.question_id
        ).first()
        verdict = auto_judge_answer(arena_question, original, ans.answer)
        if verdict is None:
            team = teams.get(ans.team_id)
            skipped.append({
                "answer_id": ans.id,
                "team_name": team.name if team else "",
                "reason": "主观题或无法自动比对，需人工评分",
            })
            continue
        result = judge_answer(db, ans.id, verdict)
        if result:
            judged.append(result)
    return judged, skipped

def compute_judgement_score(arena, arena_question, is_correct: bool, manual_score: int = None) -> int:
    """按题型规则计算该题得分（含扣分）"""
    settings = get_arena_settings(arena)
    q_type = arena_question.q_type
    if q_type == "analysis":
        # 综合分析题由主持人按要点给分；判错则不得分
        return int(manual_score or 0) if is_correct else 0
    if is_correct:
        return int(arena_question.score)
    if q_type == "rush":
        return -int(settings.get("rush_wrong_penalty", 10))
    if q_type == "overtime":
        return -int(settings.get("overtime_wrong_penalty", 15))
    if q_type == "required" and settings.get("required_wrong_penalty"):
        # 必答题开启扣分后，按本题分值扣除
        return -int(arena_question.score)
    return 0

def judge_answer(db: Session, answer_id: int, is_correct: bool = None, score: int = None):
    """主持人判定答题结果，并一次性同步队伍总分（幂等：重复判定只结算差额）"""
    answer = db.query(models.TeamAnswer).filter(models.TeamAnswer.id == answer_id).first()
    if not answer:
        return None
    team = answer.team
    arena_question = get_arena_question(db, answer.question_id)
    if not arena_question:
        return None
    original_question = db.query(models.Question).filter(
        models.Question.id == arena_question.question_id
    ).first()

    if is_correct is None:
        verdict = auto_judge_answer(arena_question, original_question, answer.answer)
        if verdict is None:
            # 无法自动判定（主观题 / 无标准答案 / 选项无文本），保持待判定状态
            return None
        is_correct = verdict
    new_score = compute_judgement_score(team.arena, arena_question, bool(is_correct), score)

    # 只结算与已入账分值的差额，避免重复判定造成分数漂移
    delta = new_score - (answer.score or 0)
    team.score = (team.score or 0) + delta
    answer.is_correct = bool(is_correct)
    answer.score = new_score
    db.commit()
    db.refresh(answer)
    return answer


def auto_judge_team_answer(db: Session, answer_id: int):
    """自动判定单条作答。

    返回 (answer, verdict)：verdict 为 True/False 表示判定结果；
    verdict 为 None 时 answer 未被修改，需要主持人人工评分。
    """
    answer = db.query(models.TeamAnswer).filter(models.TeamAnswer.id == answer_id).first()
    if not answer:
        return None, None
    arena_question = get_arena_question(db, answer.question_id)
    if not arena_question:
        return None, None
    original = db.query(models.Question).filter(
        models.Question.id == arena_question.question_id
    ).first()
    verdict = auto_judge_answer(arena_question, original, answer.answer)
    if verdict is None:
        return answer, None
    judged = judge_answer(db, answer_id, verdict)
    return judged, verdict

def update_answer_result(db: Session, answer_id: int, is_correct: bool, score: int = None):
    """兼容旧接口：主持人修改答题结果"""
    return judge_answer(db, answer_id, is_correct, score)

def player_join_arena(db: Session, join_in: schemas.ArenaPlayerJoin):
    """选手加入擂台赛"""
    team = get_arena_team(db, join_in.team_id)
    if not team:
        return None
    arena = team.arena
    if not arena or not is_invite_valid(arena):
        return None
    # 检查设备是否已在其他队伍
    existing = db.query(models.ArenaPlayer).filter(
        models.ArenaPlayer.device_id == join_in.device_id
    ).first()
    if existing:
        if existing.team_id == join_in.team_id:
            existing.last_active = datetime.utcnow()
            db.commit()
            return existing
        # 从其他队伍移除
        db.delete(existing)
    else:
        # 每个邀请码最多允许10人进入
        if count_arena_players(db, arena.id) >= 10:
            return None
    player = models.ArenaPlayer(
        team_id=join_in.team_id,
        device_id=join_in.device_id,
        player_name=join_in.player_name,
    )
    db.add(player)
    db.commit()
    db.refresh(player)
    return player

def get_team_players(db: Session, team_id: int):
    """获取队伍所有选手"""
    return db.query(models.ArenaPlayer).filter(
        models.ArenaPlayer.team_id == team_id
    ).all()

def get_arena_standings(db: Session, arena_id: int) -> schemas.ArenaStandingsOut:
    """计算排名与平局判定"""
    arena = get_arena(db, arena_id)
    if not arena:
        return schemas.ArenaStandingsOut()
    teams = get_arena_teams(db, arena_id)
    settings = get_arena_settings(arena)
    max_overtime = int(settings.get("overtime_count", 10))

    items = []
    for t in teams:
        answered = db.query(models.TeamAnswer).filter(
            models.TeamAnswer.team_id == t.id
        ).count()
        items.append({
            "id": t.id,
            "name": t.name,
            "color": t.color,
            "score": t.score or 0,
            "status": t.status,
            "rank": 0,
            "answered_count": answered,
        })
    # 排名：得分降序，同分按名称稳定排序；已淘汰队伍排在最后
    items.sort(key=lambda x: (1 if x["status"] == "eliminated" else 0, -x["score"], x["name"]))
    for idx, item in enumerate(items):
        item["rank"] = idx + 1

    active_items = [i for i in items if i["status"] != "eliminated"]
    top_score = active_items[0]["score"] if active_items else 0
    tied_ids = [i["id"] for i in active_items if i["score"] == top_score]
    tied = len(tied_ids) > 1

    # 已进行的加赛题数
    overtime_q_ids = [
        q.id for q in get_arena_questions(db, arena_id) if q.q_type == "overtime"
    ]
    overtime_count = 0
    if overtime_q_ids:
        overtime_count = db.query(models.TeamAnswer).filter(
            models.TeamAnswer.question_id.in_(overtime_q_ids)
        ).count()

    return schemas.ArenaStandingsOut(
        teams=[schemas.ArenaStandingsTeam(**i) for i in items],
        top_score=top_score,
        tied=tied,
        tied_team_ids=tied_ids,
        overtime_needed=tied and overtime_count < max_overtime,
        overtime_count=overtime_count,
        max_overtime=max_overtime,
    )

def eliminate_teams(db: Session, arena_id: int, team_ids: List[int]):
    """淘汰指定队伍（主持人手动触发）"""
    teams = get_arena_teams(db, arena_id)
    target_ids = set(team_ids or [])
    eliminated = []
    for t in teams:
        if t.id in target_ids and t.status != "eliminated":
            t.status = "eliminated"
            eliminated.append(t)
    db.commit()
    for t in eliminated:
        db.refresh(t)
    return eliminated

def eliminate_bottom_teams(db: Session, arena_id: int, count: int = None):
    """按当前排名淘汰最后N支队伍"""
    arena = get_arena(db, arena_id)
    if not arena:
        return []
    settings = get_arena_settings(arena)
    if count is None:
        count = int(settings.get("elimination_count", 1))
    standings = get_arena_standings(db, arena_id)
    active = [t for t in standings.teams if t.status != "eliminated"]
    if count <= 0 or len(active) <= count:
        return []
    bottom = [t.id for t in active[-count:]]
    return eliminate_teams(db, arena_id, bottom)

def resurrect_team(db: Session, team_id: int):
    """复活一支被淘汰的队伍"""
    team = get_arena_team(db, team_id)
    if not team:
        return None
    team.status = "resurrected"
    db.commit()
    db.refresh(team)
    return team

def get_team_by_invite_and_name(db: Session, arena_id: int, name: str):
    """按擂台赛与队名查找队伍"""
    return db.query(models.ArenaTeam).filter(
        models.ArenaTeam.arena_id == arena_id,
        models.ArenaTeam.name == name,
    ).first()

def count_arena_players(db: Session, arena_id: int) -> int:
    """统计某擂台赛已加入的选手数（用于邀请码10人上限）"""
    team_ids = [t.id for t in get_arena_teams(db, arena_id)]
    if not team_ids:
        return 0
    return db.query(models.ArenaPlayer).filter(
        models.ArenaPlayer.team_id.in_(team_ids)
    ).count()

def get_help_excluded_options(db: Session, arena_question, exclude_count: int = 2):
    """求助功能：服务端计算需要排除的错误选项（不泄露正确答案）"""
    if not arena_question:
        return []
    original = db.query(models.Question).filter(
        models.Question.id == arena_question.question_id
    ).first()
    if not original or not original.answer:
        return []
    try:
        options = json.loads(original.options or "{}")
    except (json.JSONDecodeError, TypeError):
        return []
    correct_keys = {k.strip().upper() for k in str(original.answer).split(",") if k.strip()}
    wrong_keys = [k for k in options.keys() if k.strip().upper() not in correct_keys]
    return wrong_keys[:exclude_count]

def export_arena_excel(db: Session, arena_id: int):
    """导出台词赛Excel报表（总成绩表 / 答题明细 / 统计表）"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    arena = get_arena(db, arena_id)
    if not arena:
        return None
    teams = get_arena_teams(db, arena_id)
    questions = get_arena_questions(db, arena_id)
    standings = get_arena_standings(db, arena_id)
    rank_map = {t.id: t.rank for t in standings.teams}

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="305496")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")

    def write_header(ws, row_idx, headers):
        for col, name in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col, value=name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center

    def autosize(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ===== Sheet1: 总成绩表 =====
    ws1 = wb.active
    ws1.title = "总成绩表"
    write_header(ws1, 1, ["排名", "队伍名称", "队伍颜色", "总得分", "状态"])
    row = 2
    for t in sorted(teams, key=lambda x: rank_map.get(x.id, 999)):
        ws1.cell(row=row, column=1, value=rank_map.get(t.id, row - 1))
        ws1.cell(row=row, column=2, value=t.name)
        c = ws1.cell(row=row, column=3, value=t.color)
        c.fill = PatternFill("solid", fgColor=(t.color or "#409EFF").lstrip("#"))
        ws1.cell(row=row, column=4, value=t.score or 0)
        status_map = {"active": "参赛中", "eliminated": "已淘汰", "resurrected": "已复活"}
        ws1.cell(row=row, column=5, value=status_map.get(t.status, t.status))
        row += 1
    autosize(ws1, [8, 20, 12, 12, 12])

    # ===== Sheet2: 答题明细 =====
    ws2 = wb.create_sheet("答题明细")
    write_header(ws2, 1, ["队伍", "题号", "题型", "题目内容", "正确答案", "队伍答案",
                          "是否正确", "得分", "用时(秒)", "求助", "提交时间"])
    type_map = {"required": "必答题", "rush": "抢答题", "analysis": "综合分析题", "overtime": "加赛题"}
    row = 2
    for t in teams:
        for q in questions:
            rec = db.query(models.TeamAnswer).filter(
                models.TeamAnswer.team_id == t.id,
                models.TeamAnswer.question_id == q.id,
            ).first()
            original = db.query(models.Question).filter(
                models.Question.id == q.question_id
            ).first()
            content = ""
            if original and original.content:
                content = original.content.replace("<p>", "").replace("</p>", "")[:120]
            correct_answer = original.answer if original else ""
            if q.q_type == "analysis":
                correct_answer = "（主持人评分）"
            ws2.cell(row=row, column=1, value=t.name)
            ws2.cell(row=row, column=2, value=q.order_num)
            ws2.cell(row=row, column=3, value=type_map.get(q.q_type, q.q_type))
            ws2.cell(row=row, column=4, value=content)
            ws2.cell(row=row, column=5, value=correct_answer)
            ws2.cell(row=row, column=6, value=rec.answer if rec else "")
            if rec is None:
                verdict = "未作答"
            elif rec.is_correct is None:
                verdict = "待判定"
            else:
                verdict = "正确" if rec.is_correct else "错误"
            ws2.cell(row=row, column=7, value=verdict)
            ws2.cell(row=row, column=8, value=rec.score if rec else 0)
            ws2.cell(row=row, column=9, value=rec.time_used if rec else 0)
            ws2.cell(row=row, column=10, value=rec.help_used if rec else "")
            ws2.cell(row=row, column=11,
                     value=rec.submitted_at.strftime("%Y-%m-%d %H:%M:%S") if rec else "")
            row += 1
    autosize(ws2, [16, 8, 12, 50, 14, 24, 10, 8, 10, 12, 20])

    # ===== Sheet3: 统计表 =====
    ws3 = wb.create_sheet("统计表")
    write_header(ws3, 1, ["题型", "题目数", "作答数", "正确数", "正确率", "平均用时(秒)"])
    row = 2
    for q_type, label in [("required", "必答题"), ("rush", "抢答题"),
                          ("analysis", "综合分析题"), ("overtime", "加赛题")]:
        qs = [q for q in questions if q.q_type == q_type]
        q_ids = [q.id for q in qs]
        if not q_ids:
            ws3.cell(row=row, column=1, value=label)
            for c in range(2, 7):
                ws3.cell(row=row, column=c, value=0 if c != 5 else "-")
            row += 1
            continue
        records = db.query(models.TeamAnswer).filter(
            models.TeamAnswer.question_id.in_(q_ids)
        ).all()
        total = len(qs) * len(teams) if teams else 0
        answered = len(records)
        correct = len([r for r in records if r.is_correct is True])
        rate = f"{(correct / answered * 100):.1f}%" if answered else "-"
        avg_time = f"{sum(r.time_used or 0 for r in records) / answered:.1f}" if answered else "-"
        ws3.cell(row=row, column=1, value=label)
        ws3.cell(row=row, column=2, value=len(qs))
        ws3.cell(row=row, column=3, value=answered)
        ws3.cell(row=row, column=4, value=correct)
        ws3.cell(row=row, column=5, value=rate)
        ws3.cell(row=row, column=6, value=avg_time)
        row += 1
    autosize(ws3, [16, 10, 10, 10, 10, 14])

    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def log_arena_action(db: Session, arena_id: int, action: str, actor_type: str, actor_id: int = None, actor_name: str = "", detail: str = ""):
    """记录操作日志"""
    log = models.ArenaLog(
        arena_id=arena_id,
        action=action,
        actor_type=actor_type,
        actor_id=actor_id,
        actor_name=actor_name,
        detail=detail,
    )
    db.add(log)
    db.commit()
    db.refresh(log)
    return log

def get_arena_logs(db: Session, arena_id: int, limit=100):
    """获取擂台赛操作日志"""
    return db.query(models.ArenaLog).filter(
        models.ArenaLog.arena_id == arena_id
    ).order_by(models.ArenaLog.timestamp.desc()).limit(limit).all()

def create_arena_round(db: Session, arena_id: int, round_in: schemas.ArenaRoundCreate):
    """创建轮次"""
    arena = get_arena(db, arena_id)
    if not arena:
        return None
    round_record = models.ArenaRound(
        arena_id=arena_id,
        round_num=round_in.round_num,
        round_type=round_in.round_type,
        status="pending",
    )
    db.add(round_record)
    db.commit()
    db.refresh(round_record)
    return round_record

def get_arena_rounds(db: Session, arena_id: int):
    """获取擂台赛所有轮次"""
    return db.query(models.ArenaRound).filter(
        models.ArenaRound.arena_id == arena_id
    ).order_by(models.ArenaRound.round_num).all()

def update_arena_round(db: Session, round_id: int, status: str):
    """更新轮次状态"""
    round_record = db.query(models.ArenaRound).filter(models.ArenaRound.id == round_id).first()
    if not round_record:
        return None
    round_record.status = status
    if status == "active":
        round_record.start_time = datetime.utcnow()
    elif status == "finished":
        round_record.end_time = datetime.utcnow()
    db.commit()
    db.refresh(round_record)
    return round_record

def export_arena_answers(db: Session, arena_id: int) -> List[schemas.ArenaAnswerExport]:
    """导出台词答案"""
    arena = get_arena(db, arena_id)
    if not arena:
        return []
    teams = get_arena_teams(db, arena_id)
    questions = get_arena_questions(db, arena_id)
    result = []
    for team in teams:
        answers = []
        for aq in questions:
            answer_record = db.query(models.TeamAnswer).filter(
                models.TeamAnswer.team_id == team.id,
                models.TeamAnswer.question_id == aq.id,
            ).first()
            original_q = db.query(models.Question).filter(
                models.Question.id == aq.question_id
            ).first()
            answers.append({
                "question_index": aq.order_num,
                "question_type": aq.q_type,
                "question_content": original_q.content[:100] + "..." if original_q and len(original_q.content) > 100 else (original_q.content if original_q else ""),
                "correct_answer": original_q.answer if original_q else "",
                "team_answer": answer_record.answer if answer_record else "",
                "is_correct": answer_record.is_correct if answer_record else None,
                "score": answer_record.score if answer_record else 0,
            })
        result.append(schemas.ArenaAnswerExport(
            team_name=team.name,
            team_color=team.color,
            total_score=team.score,
            answers=answers,
        ))
    return result


def _exam_answer_correct(question, user_answer: str) -> bool:
    """按考试交卷规则判断某题对错（多选按集合比较，其余按去空格字符串比较）。"""
    given = (user_answer or "").strip()
    if question.type == "multiple":
        correct = {x.strip() for x in (question.answer or "").split(",") if x.strip()}
        return bool(given) and {x.strip() for x in given.split(",") if x.strip()} == correct
    if question.type == "fill":
        return fill_answer_matches(given, question.answer)
    return given == (question.answer or "").strip()


def export_exam_excel(db: Session, exam_id: int):
    """导出普通考试成绩 Excel（总成绩表 / 答题明细 / 统计表）"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    exam = get_exam(db, exam_id)
    if not exam:
        return None

    records = get_exam_records(db, exam_id)
    # 随机卷：每个考生抽到的题不一样，逐题明细按本人的卷子走，
    # 统计表用所有考生卷子的并集（没抽到的题不会出现）。
    per_user_questions = {}
    if getattr(exam, "random_mode", False):
        papers = db.query(models.ExamPaper).filter(
            models.ExamPaper.exam_id == exam.id).all()
        union_ids = []
        seen_ids = set()
        for p in papers:
            try:
                ids = json.loads(p.question_ids or "[]")
            except (TypeError, ValueError):
                ids = []
            for qid in ids:
                if qid not in seen_ids:
                    seen_ids.add(qid)
                    union_ids.append(qid)
        rows = db.query(models.Question).filter(
            models.Question.id.in_(union_ids)).all() if union_ids else []
        by_id = {q.id: q for q in rows}
        union_questions = [by_id[i] for i in union_ids if i in by_id]
        for p in papers:
            try:
                ids = json.loads(p.question_ids or "[]")
            except (TypeError, ValueError):
                ids = []
            per_user_questions[(p.user_id, p.attempt_no)] = [
                by_id[i] for i in ids if i in by_id]
        # 成绩 -> 对应轮次（同一考生补考时，每轮卷子不同）
        attempts = db.query(models.ExamAttempt).filter(
            models.ExamAttempt.exam_id == exam.id).all()
        for rec in records:
            cands = [a for a in attempts
                     if a.user_id == rec.user_id and a.submitted_at is not None
                     and a.submitted_at <= (rec.submitted_at or a.submitted_at)]
            if cands:
                best = max(cands, key=lambda a: a.attempt_no)
                per_user_questions[("record", rec.id)] = per_user_questions.get(
                    (rec.user_id, best.attempt_no), [])
        questions = sorted(union_questions, key=lambda q: q.id)
    else:
        questions = sorted(exam.questions, key=lambda q: q.id)
    score_map = {
        "single": exam.single_score or 10,
        "multiple": exam.multiple_score or 20,
        "judgment": exam.judgment_score or 5,
    }
    type_label = {"single": "单选题", "multiple": "多选题", "judgment": "判断题", "fill": "填空题"}

    # 解析每份答卷的作答明细
    parsed = []
    for r in records:
        try:
            answers = json.loads(r.answers or "{}")
        except (json.JSONDecodeError, TypeError):
            answers = {}
        if not isinstance(answers, dict):
            answers = {}
        parsed.append((r, answers))

    def percent_of(r):
        return (r.score / r.total_score * 100) if r.total_score else 0

    # 按百分制排序（并列按提交时间先后）
    parsed.sort(key=lambda item: (-percent_of(item[0]), item[0].submitted_at or datetime.min))
    rank_map = {}
    last_pct = None
    last_rank = 0
    for idx, (r, _) in enumerate(parsed, start=1):
        pct = percent_of(r)
        if last_pct is not None and pct == last_pct:
            rank_map[r.id] = last_rank
        else:
            rank_map[r.id] = idx
            last_rank = idx
        last_pct = pct

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="305496")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")

    def write_header(ws, row_idx, headers):
        for col, name in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col, value=name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center

    def autosize(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ===== Sheet1: 总成绩表 =====
    ws1 = wb.active
    ws1.title = "总成绩表"
    write_header(ws1, 1, ["排名", "工号", "姓名", "部门", "得分", "总分", "百分制", "是否及格", "提交时间"])
    row = 2
    for r, _ in parsed:
        pct = percent_of(r)
        passed = pct >= (exam.pass_score or 0)
        ws1.cell(row=row, column=1, value=rank_map.get(r.id, row - 1))
        ws1.cell(row=row, column=2, value=r.user.employee_id if r.user else "")
        ws1.cell(row=row, column=3, value=r.user.name if r.user else "")
        ws1.cell(row=row, column=4, value=r.user.dept if r.user else "")
        ws1.cell(row=row, column=5, value=r.score or 0)
        ws1.cell(row=row, column=6, value=r.total_score or 0)
        ws1.cell(row=row, column=7, value=round(pct, 1))
        ws1.cell(row=row, column=8, value="及格" if passed else "不及格")
        ws1.cell(row=row, column=9,
                 value=r.submitted_at.strftime("%Y-%m-%d %H:%M:%S") if r.submitted_at else "")
        row += 1
    autosize(ws1, [8, 12, 12, 16, 10, 10, 10, 12, 22])

    # ===== Sheet2: 答题明细 =====
    ws2 = wb.create_sheet("答题明细")
    write_header(ws2, 1, ["工号", "姓名", "题号", "题型", "题目内容", "正确答案", "考生答案", "是否正确", "得分"])
    row = 2
    for r, answers in parsed:
        # 随机卷按该考生自己那份卷子出题，普通卷用整卷题目
        row_questions = questions
        if per_user_questions:
            row_questions = per_user_questions.get(
                ("record", r.id)) or questions
        for idx, q in enumerate(row_questions, start=1):
            user_answer = answers.get(str(q.id), answers.get(q.id, ""))
            correct = _exam_answer_correct(q, user_answer)
            content = (q.content or "").replace("<p>", "").replace("</p>", "")[:120]
            ws2.cell(row=row, column=1, value=r.user.employee_id if r.user else "")
            ws2.cell(row=row, column=2, value=r.user.name if r.user else "")
            ws2.cell(row=row, column=3, value=idx)
            ws2.cell(row=row, column=4, value=type_label.get(q.type, q.type))
            ws2.cell(row=row, column=5, value=content)
            ws2.cell(row=row, column=6, value=q.answer or "")
            ws2.cell(row=row, column=7, value=user_answer or "")
            ws2.cell(row=row, column=8, value="正确" if correct else "错误")
            ws2.cell(row=row, column=9, value=score_map.get(q.type, 10) if correct else 0)
            row += 1
    autosize(ws2, [12, 12, 8, 12, 50, 14, 24, 10, 8])

    # ===== Sheet3: 统计表 =====
    ws3 = wb.create_sheet("统计表")
    write_header(ws3, 1, ["题号", "题型", "题目内容", "正确答案", "作答人数", "答对人数", "正确率"])
    row = 2
    for idx, q in enumerate(questions, start=1):
        answered = 0
        correct_count = 0
        for r, answers in parsed:
            user_answer = answers.get(str(q.id), answers.get(q.id, ""))
            if not (user_answer or "").strip():
                continue
            answered += 1
            if _exam_answer_correct(q, user_answer):
                correct_count += 1
        rate = f"{(correct_count / answered * 100):.1f}%" if answered else "-"
        ws3.cell(row=row, column=1, value=idx)
        ws3.cell(row=row, column=2, value=type_label.get(q.type, q.type))
        ws3.cell(row=row, column=3, value=(q.content or "").replace("<p>", "").replace("</p>", "")[:80])
        ws3.cell(row=row, column=4, value=q.answer or "")
        ws3.cell(row=row, column=5, value=answered)
        ws3.cell(row=row, column=6, value=correct_count)
        ws3.cell(row=row, column=7, value=rate)
        row += 1

    # 整体统计追加在逐题统计下方
    row += 1
    summary_title = row
    ws3.cell(row=row, column=1, value="整体统计").font = Font(bold=True)
    row += 1
    total = len(parsed)
    pcts = [percent_of(r) for r, _ in parsed]
    passed = len([p for p in pcts if p >= (exam.pass_score or 0)])
    summary = [
        ("参考人数", total),
        ("平均百分制", round(sum(pcts) / total, 1) if total else 0),
        ("最高百分制", round(max(pcts), 1) if pcts else 0),
        ("最低百分制", round(min(pcts), 1) if pcts else 0),
        ("及格人数", passed),
        ("及格率", f"{(passed / total * 100):.1f}%" if total else "-"),
        ("及格线", exam.pass_score or 0),
    ]
    for label, value in summary:
        ws3.cell(row=row, column=1, value=label)
        ws3.cell(row=row, column=2, value=value)
        row += 1
    autosize(ws3, [8, 12, 50, 14, 12, 12, 10])

    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
